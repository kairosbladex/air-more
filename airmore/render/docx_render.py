"""DOCX 渲染：A4 版式，中文字体（宋体正文、黑体标题），页脚统一标注。

报头仿政府日报三行式（红色题名 / 期号 / 日期，下加红色横线）。
正文为块驱动渲染（section / caption / paragraph / note / table / image 等），
A 类为政府日报版式；B/G 类为摘要页 + 事实证据 + 附录版式，报头统一。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from airmore import DISCLAIMER, THRESHOLD_DISCLAIMER

DARK = RGBColor(0x1F, 0x1F, 0x1F)
GRAY = RGBColor(0x66, 0x66, 0x66)
ACCENT = RGBColor(0xB0, 0x30, 0x30)      # 风险强调色
HEADER_RED = RGBColor(0xC0, 0x00, 0x00)  # 报头红
RANK_COLORS = {"yellow": RGBColor(0xFF, 0xC0, 0x00), "red": RGBColor(0xFF, 0x00, 0x00)}


def _set_eastasia(style_or_run, ascii_font: str, eastasia_font: str) -> None:
    style_or_run.font.name = ascii_font
    rPr = style_or_run.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), eastasia_font)


def _setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.space_before = Pt(0)
    _set_eastasia(normal, "Times New Roman", "宋体")
    for name, size in [("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)]:
        st = doc.styles[name]
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = DARK
        st.paragraph_format.space_before = Pt(4)
        st.paragraph_format.space_after = Pt(2)
        _set_eastasia(st, "Arial", "黑体")


def _setup_page(doc: Document, disclaimer: str = DISCLAIMER) -> None:
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)  # A4
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.left_margin = sec.right_margin = Cm(2.2)
    # 页脚：数据声明 + 页码
    footer_p = sec.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run(f"{disclaimer}｜第 ")
    run.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer_p._p.append(fld)
    run2 = footer_p.add_run(" 页")
    run2.font.size = Pt(8)


def new_document(disclaimer: str = DISCLAIMER) -> Document:
    doc = Document()
    _setup_styles(doc)
    _setup_page(doc, disclaimer)
    return doc


# ── 报头 ──────────────────────────────────────────────────────────────

def _render_header(doc: Document, header: dict) -> None:
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run(header["line1"])
    r1.font.size = Pt(header.get("line1_size", 24))
    r1.font.bold = True
    r1.font.color.rgb = HEADER_RED
    _set_eastasia(r1, "Times New Roman", "黑体")

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(header["line2"])
    r2.font.size = Pt(15)
    r2.font.bold = True

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r3 = p3.add_run(header["line3"])
    r3.font.size = Pt(15)
    r3.font.bold = True

    # 红色横线（红头文件样式）
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(4)
    pPr = rule._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:color"), "C00000")
    pBdr.append(bottom)
    pPr.append(pBdr)


# ── 块渲染 ────────────────────────────────────────────────────────────

def _apply_break(paragraph) -> None:
    paragraph.paragraph_format.page_break_before = True


def _data_table(doc: Document, header: list[str], rows: list[list[str]],
                font_size: float = 9, col_widths: list[float] | None = None,
                row_colors: dict[int, str] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.text = h
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(font_size)
    for i, row in enumerate(rows):
        color = RANK_COLORS.get((row_colors or {}).get(i))
        cells = table.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = str(v)
            for p in cells[j].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(font_size)
                    if color is not None:
                        r.font.color.rgb = color
    if col_widths:
        table.autofit = False  # 固定布局 + 网格列宽，LibreOffice 才认
        for j, w in enumerate(col_widths):
            table.columns[j].width = Cm(w)
        for row in table.rows:
            for j, w in enumerate(col_widths):
                row.cells[j].width = Cm(w)


def _info_table(doc: Document, rows: list[tuple[str, str]],
                font_size: float = 10.5) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        c0.text, c1.text = k, v
        for p in c0.paragraphs:
            for r in p.runs:
                r.font.bold = True
        for c in (c0, c1):
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(font_size)
        c0.width, c1.width = Cm(4.2), Cm(12.2)


def _picture(doc: Document, path: str | Path, width_cm: float = 13.0,
             caption: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    if caption:
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(2)
        for r in cap.runs:
            r.font.size = Pt(8)
            r.font.color.rgb = GRAY


def _picture_row(doc: Document, items: list[tuple[str | Path, float]]) -> None:
    """多张图并排放在同一段（总宽度不超过版心 16.6 cm）。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    for path, width_cm in items:
        p.add_run().add_picture(str(path), width=Cm(width_cm))


def _render_block(doc: Document, block: dict, state: dict) -> None:
    btype = block["type"]
    need_break = state.pop("page_break", False)

    if btype == "page_break":
        state["page_break"] = True
        return
    if btype == "section":
        h = doc.add_heading(block["text"], level=1)
        if need_break:
            _apply_break(h)
        return
    if btype == "heading":
        h = doc.add_heading(block["text"], level=2)
        if need_break:
            _apply_break(h)
        return
    if btype == "caption":
        p = doc.add_paragraph(block["text"])
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if need_break:
            _apply_break(p)
        for r in p.runs:
            r.font.size = Pt(11)
        return
    if btype == "paragraph":
        # 文本内的换行拆分为多个段落（Word 不认 run 内 \n）
        for i, seg in enumerate(block["text"].split("\n")):
            p = doc.add_paragraph(seg)
            if need_break and i == 0:
                _apply_break(p)
        return
    if btype == "note":
        p = doc.add_paragraph(block["text"])
        if need_break:
            _apply_break(p)
        for r in p.runs:
            r.font.size = Pt(9)
            r.font.color.rgb = GRAY
        return
    if btype == "conclusions":
        if need_break:
            state["page_break"] = True
        for c in block["items"]:
            p = doc.add_paragraph(style="List Number")
            run = p.add_run(c)
            if any(k in c for k in ("优先核查", "建议复核")):
                run.font.color.rgb = ACCENT
        return
    if btype == "bullets":
        if need_break:
            state["page_break"] = True
        for a in block["items"]:
            doc.add_paragraph(a, style="List Bullet")
        return
    if btype == "meta_table":
        if need_break:
            # 表格无法直接分页，前置一个最小段落承接分页
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.font.size = Pt(1)
            _apply_break(p)
        _info_table(doc, block["rows"], font_size=block.get("font_size", 10.5))
        return
    if btype == "table":
        if need_break:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.font.size = Pt(1)
            _apply_break(p)
        _data_table(doc, block["header"], block["rows"],
                    font_size=block.get("font_size", 9),
                    col_widths=block.get("col_widths"),
                    row_colors=block.get("row_colors"))
        return
    if btype == "image":
        _picture(doc, block["path"], width_cm=block.get("width_cm", 13.0),
                 caption=block.get("caption"))
        return
    if btype == "image_row":
        _picture_row(doc, [(it["path"], it["width_cm"])
                           for it in block["images"]])
        if block.get("caption"):
            cap = doc.add_paragraph(block["caption"])
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in cap.runs:
                r.font.size = Pt(8)
                r.font.color.rgb = GRAY
        return
    raise ValueError(f"未知块类型: {btype}")


def render_report_docx(spec: dict, out_path: str | Path) -> Path:
    """按结构化 spec 渲染 DOCX。PDF 由同一 DOCX 转换，保证两版事实一致。"""
    doc = new_document(spec.get("disclaimer", DISCLAIMER))
    _render_header(doc, spec["header"])
    state: dict = {}
    for block in spec["body"]:
        _render_block(doc, block, state)
    out_path = Path(out_path)
    doc.save(out_path)
    return out_path


# ── 共用块构造 ────────────────────────────────────────────────────────

def appendix_blocks(qc_rows: list[list[str]], rule_lines: list[str],
                    source_lines: list[str],
                    title: str = "建议与附录",
                    disclaimer: str = DISCLAIMER,
                    disclaimer_suffix: str = "虚构机库与示意航线，不指向真实主体。") -> list[dict]:
    """数据质量 + 规则/来源/声明（B/G 附录，A 附注共用）。"""
    blocks = [
        {"type": "section", "text": title},
        {"type": "heading", "text": "数据质量"},
        {"type": "table", "header": ["检查项", "数量", "说明"], "rows": qc_rows,
         "font_size": 8, "col_widths": [4.5, 1.4, 10.1]},
        {"type": "heading", "text": "规则、来源与声明"},
    ]
    for line in rule_lines + source_lines + [
            THRESHOLD_DISCLAIMER + "；" + disclaimer + "。" + disclaimer_suffix]:
        blocks.append({"type": "note", "text": line})
    return blocks
