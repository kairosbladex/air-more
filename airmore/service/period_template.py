"""Period-specific reports retaining the reference document's base page/style source."""
from datetime import datetime
from hashlib import sha256

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .models import LABELS, METRICS, TZ, survey_value_text
from .strict_template import TEMPLATE
from .advice import HORIZONS

VERSION = 'period-survey-report/v1'


def render_document(stats, analyses, composite, cfg, output, runtime):
    doc = Document(TEMPLATE)
    body = doc._element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)
    # Keep source document styles and page dimensions; use flowing business content.
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style = doc.styles[name]
        style.font.name = runtime['font_family']
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), runtime.get('docx_font_family', runtime['font_family']))
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt({'Normal': 10.5, 'Title': 20, 'Heading 1': 13, 'Heading 2': 11.5}[name])
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.left_indent = Cm(0)
        style.paragraph_format.right_indent = Cm(0)
        style.paragraph_format.keep_with_next = name != 'Normal'
        snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0')
        style.element.get_or_add_pPr().append(snap)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.space_before = Pt(6 if name != 'Normal' else 0)
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    req = stats['request']
    kind = req['report_type']
    title = {'DAY': '日报', 'WEEK': '周报', 'MONTH': '月报'}[kind]
    cutoff = datetime.fromisoformat(req['as_of']).astimezone(TZ) if req.get('as_of') else None
    doc.add_paragraph('空气质量巡测研判' + title, 'Title')
    window = req['start_date']+' 至 '+req['end_date']
    if cutoff and cutoff.date().isoformat() <= req['end_date']:
        window += f'  截至 {cutoff:%Y-%m-%d %H:%M}  阶段报告'
    doc.add_paragraph(window)
    if kind == 'WEEK':
        from datetime import date
        a, b = date.fromisoformat(req['start_date']), date.fromisoformat(req['end_date'])
        if a.weekday() != 0 or (b-a).days != 6:
            doc.add_paragraph('本期为选定时段周报，未覆盖完整自然周。')
    if kind == 'MONTH':
        import calendar
        from datetime import date
        a, b = date.fromisoformat(req['start_date']), date.fromisoformat(req['end_date'])
        if a.day != 1 or a.month != b.month or b.day != calendar.monthrange(b.year, b.month)[1]:
            doc.add_paragraph('本期为月度阶段分析，结论仅覆盖已观测日期。')
    if stats['data_kind'] == 'synthetic':
        doc.add_paragraph('合成样本用于计算测试，不用于业务决策。')
    historical_stream = stats['data_kind'] == 'historical_sensor_replay'
    if historical_stream:
        doc.add_paragraph('历史传感器数据复算。原始采集记录仍在，但对应任务关联缺失，无法确认历史机库归属；本报告按项目和传感器归集，用于报告效果验证。工作建议按所选历史报告期生成。')
        doc.add_paragraph('单位按本次核验的同项目产品配置及单位字典解释，不代表已追溯确认历史配置。气体保留原始ppm或ppb，不作温压换算或正式AQI评价。')

    def table(headers, rows):
        t = doc.add_table(rows=1, cols=len(headers))
        t.autofit = False
        widths = [16.5/len(headers)] * len(headers)
        for c, w in zip(t.columns, widths):
            c.width = Cm(w)
        for c, label in zip(t.rows[0].cells, headers):
            c.text = label
        for values in rows:
            for c, value in zip(t.add_row().cells, values):
                c.text = str(value)
        for n, row in enumerate(t.rows):
            pr = row._tr.get_or_add_trPr()
            pr.append(OxmlElement('w:cantSplit'))
            if n == 0:
                pr.append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tc = cell._tc.get_or_add_tcPr()
                borders = OxmlElement('w:tcBorders')
                for edge in ('top','left','bottom','right'):
                    e = OxmlElement('w:'+edge)
                    for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(qn('w:'+k),v)
                    borders.append(e)
                tc.append(borders)
                if n == 0:
                    shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E8EEF2');tc.append(shading)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if len(t.rows) <= 8:
                        p.paragraph_format.keep_with_next = n < len(t.rows) - 1
                    p.paragraph_format.first_line_indent = Cm(0)
                    snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0'); p._p.get_or_add_pPr().append(snap)
                    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(2 if len(t.rows) > 8 else 4)
                    for run in p.runs:
                        run.font.size = Pt(9)
        return t

    for hid, h in stats['hangars'].items():
        c, r = h['current'], h['relative']
        doc.add_heading(h['name'], level=1)
        def value_text(v):
            return survey_value_text(v['value'], v['unit'])
        levels = [f"{LABELS[m]} {value_text(v)} {v['unit']}" for m,v in c['survey'].items() if v['value'] is not None]
        sample_label = '归集采样' if historical_stream else '有效关联采样'
        doc.add_paragraph(f"本期有观测{c['observed_days']}日，区间共{c['days']}日，{sample_label}{c['samples']}条。" + ('可用指标为'+ '，'.join(levels)+'。' if levels else '浓度统计所需的有效值或单位信息不足。') + '巡测结果仅代表实际采样覆盖。')
        doc.add_heading(HORIZONS[kind], level=1)
        analysis = analyses[hid]
        if analysis.metadata.get('engine') != 'ollama':
            doc.add_paragraph('模型研判待生成。以下为显式模板模式的资料核查提示。')
        for paragraph in analysis.sections['advice'].split('\n'):
            doc.add_paragraph(paragraph)
        doc.add_heading('浓度与观测覆盖', level=1)
        table(['因子','巡测统计值','单位','有效日期'], [[LABELS[m],value_text(c['survey'][m]), c['survey'][m]['unit'] or '待核验',f"{c['survey'][m]['days']}/{c['days']}"] for m in METRICS])
        doc.add_paragraph('日内先对有效分钟等权统计，再对有数据日期等权；斜线表示不可计算，不表示零。')
        doc.add_heading('可比历史与指标分解', level=1)
        if r['current']:
            v = r['current']
            doc.add_paragraph(f"颗粒物巡测相对指数为{v['value']:.2f}，历史基准为100。可用可比观测层{r['usable_strata']}/{r['current_strata']}；不作跨机库排名。")
            if r['coverage'] < 1:
                doc.add_paragraph(f"可比层覆盖率仅{r['coverage']:.2%}，上述指数只描述这部分观测层，不能概括整期或整个区域的升降。")
            table(['因子','指数项','因子指数占比'],[[LABELS[m],f"{v['terms'][m]:.2f}",f"{v['ratios'][m]*100:.1f}%" if v['ratios'] else '/'] for m in r['panel']])
            change = r['changes']
            doc.add_paragraph(f"相对{change['basis']}，同范围指数由{change['previous_value']:.2f}变为{change['current_value']:.2f}，净变化{change['net']:+.2f}个指数点。")
            table(['因子','变化指数点','正向变化贡献占比'],[[LABELS[m],f"{change['terms'][m]:+.2f}", f"{change['positive_ratios'][m]*100:.1f}%" if change['positive_ratios'] else '/'] for m in r['panel']])
            if not change['positive_ratios']:doc.add_paragraph('本期未出现正向变化项。')
            if any(s['baseline_days'] < 14 for s in r['strata']):doc.add_paragraph('部分观测层只有七至十三个历史有效日，本期采用短期参考。')
        else:
            doc.add_paragraph(f"本期保留实测浓度和覆盖分析。颗粒物历史对比尚缺：{r['reason']}。")
        if r['current']:
            doc.add_paragraph('相对指数固定使用两项颗粒物因子等权计算，以报告期前三十天内可比观测的日中位数作基准；可比条件包括同传感器、同空间网格、高度层和时段。占比是指标分解，不是污染源或气象贡献。')
        if kind != 'DAY':
            doc.add_heading('逐日覆盖与复核重点', level=1)
            table(['日期','有效采样','可用浓度因子'],[[d,v['samples'],sum(x['value'] is not None for x in v['survey'].values())] for d,v in h['days'].items() if c['start'] <= d <= c['end']])
            repetitions = [x['text'] for x in h['evidence_cards'] if x['kind']=='repeat']
            for text in repetitions[:6]:doc.add_paragraph(text)
            if not repetitions:doc.add_paragraph('当前证据未形成可比历史下的重复偏高清单；不能将缺测日期判为正常，也不据此推断长期规律。')
    note = doc.add_paragraph('资料说明：本报告用于巡测复核与观测安排。完整日评价条件未满足时不计算正式日AQI；缺少同期气象或正式预报资料时，不推断气象成因及未来扩散条件。')
    note.paragraph_format.space_after = Pt(3)
    for run in note.runs: run.font.size = Pt(9)
    p = stats.get('provenance', {})
    if p:
        note = doc.add_paragraph(f"来源：业务采集表，按采集时间统计；原始本期记录{p.get('raw_current_rows', '/')}条。当前设备关联不证明历史绑定。")
        for run in note.runs: run.font.size = Pt(9)
    target = output/'report.docx'
    doc.save(target)
    loaded = Document(target)
    if len(loaded.tables) != len(doc.tables):raise ValueError('Report table readback failed')
    return {'template_version': VERSION, 'template_sha256': sha256(TEMPLATE.read_bytes()).hexdigest(),
            'readback': 'passed', 'report_type': kind, 'runtime': {k:v for k,v in runtime.items() if k!='env'}}
