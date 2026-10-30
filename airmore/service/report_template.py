"""Template-derived air-quality reports sized to the available measurements."""
from hashlib import sha256
import io
from zipfile import ZipFile, ZIP_DEFLATED

import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .models import METRICS, survey_value_text
from .strict_template import TEMPLATE, png
from .report_facts import series, fact_pack, LABELS

VERSION='air-quality-report-layout/v4'
COLORS={'pm10':'#264F70','pm2_5':'#3D8D91','so2':'#627F9A','no2':'#B58849','co':'#795D80','o3':'#5A8764'}


def trend_image(h,runtime):
    points,basis=series(h)
    active=[i for i,(_,v) in enumerate(points) if any(x.get('value') is not None for x in v.values())]
    if not active:return None
    if h['report_type']=='DAY':
        points=points[max(0,active[0]-1):min(len(points),active[-1]+2)]
    groups={}
    for m in ('pm10','pm2_5','so2','no2','co','o3'):
        units={v[m]['unit'] for _,v in points if v.get(m,{}).get('value') is not None}
        if len(units)==1:groups.setdefault(next(iter(units)),[]).append(m)
    with plt.rc_context({'font.family':runtime['font_family'],'axes.unicode_minus':False}):
        fig,axes=plt.subplots(len(groups),1,figsize=(7.6,1.65*len(groups)+.35),squeeze=False)
        for ax,(unit,metrics) in zip(axes[:,0],groups.items()):
            ax.set_facecolor('#F6F8FA')
            for m in metrics:
                vals=[v.get(m,{}).get('value') for _,v in points]
                ax.plot(range(len(points)),[float('nan') if v is None else v for v in vals],label=LABELS[m],color=COLORS[m],linewidth=1.8,marker='o',markersize=3)
            ticks=sorted(set([0,len(points)-1,*range(0,len(points),max(1,len(points)//7))]))
            if len(points)>14 and len(ticks)>2 and ticks[-1]-ticks[-2]<max(2,len(points)//10):
                ticks.pop(-2)
            labels=[points[i][0] if len(points[i][0])==5 else points[i][0][5:] for i in ticks]
            ax.set_xticks(ticks,labels,fontsize=9)
            ax.tick_params(axis='both',length=0,labelsize=9,pad=6,colors='#444444')
            ax.set_ylabel(unit,fontsize=9,color='#444444')
            ax.set_ylim(bottom=0)
            ax.grid(axis='y',color='#DEE4E8',linewidth=.6)
            for spine in ax.spines.values():spine.set_visible(False)
            ax.legend(loc='lower left',bbox_to_anchor=(0,1.01),ncol=len(metrics),frameon=False,fontsize=9,borderaxespad=0)
        fig.subplots_adjust(left=.09,right=.98,top=.9,bottom=.13,hspace=.9)
        return png(fig)


def render_document(stats,analyses,composite,cfg,output,runtime):
    doc=Document(TEMPLATE)
    for child in list(doc._element.body):
        if child.tag!=qn('w:sectPr'):doc._element.body.remove(child)
    section=doc.sections[0]
    section.top_margin=section.bottom_margin=Cm(1.8)
    section.left_margin=section.right_margin=Cm(2.2)
    font=runtime.get('docx_font_family',runtime['font_family'])

    def format_run(run,size=11,bold=False,color='000000'):
        run.font.name='Arial';run.font.size=Pt(size);run.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),font)
        sz=OxmlElement('w:szCs');sz.set(qn('w:val'),str(int(size*2)));run._element.get_or_add_rPr().append(sz)
    def para(content='',size=11,bold=False,before=0,after=6,keep=False,center=False):
        p=doc.add_paragraph(style='Normal')
        pf=p.paragraph_format
        pf.line_spacing=1.15 if size<=9 or size>=20 else 1.3;pf.space_before=Pt(before);pf.space_after=Pt(after)
        pf.first_line_indent=Pt(0);pf.left_indent=pf.right_indent=Pt(0)
        pf.keep_with_next=keep;pf.widow_control=True
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
        pp=p._p.get_or_add_pPr()
        pp.get_or_add_ind().set(qn('w:firstLineChars'),'0')
        for key in ('beforeLines','afterLines'):pp.get_or_add_spacing().set(qn('w:'+key),'0')
        snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pp.append(snap)
        format_run(p.add_run(content),size,bold)
        return p
    def heading(s):return para(s,15,True,before=8,after=7,keep=True)
    def note(s):return para(s,9,after=5)
    def table(headers,rows,widths):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        for col,w in zip(t.columns,widths):col.width=Cm(w)
        for values in [headers]+rows:
            row=t.rows[0] if values is headers else t.add_row()
            n=len(t.rows)-1 if values is not headers else 0
            pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
            if values is headers:pr.append(OxmlElement('w:tblHeader'))
            for cell,w,value in zip(row.cells,widths,values):
                cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text=str(value);tc=cell._tc.get_or_add_tcPr()
                borders=OxmlElement('w:tcBorders')
                for edge in ('top','left','bottom','right'):
                    node=OxmlElement('w:'+edge)
                    for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:node.set(qn('w:'+k),v)
                    borders.append(node)
                tc.append(borders)
                if values is headers or n%2==0:
                    shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E7E9EB' if values is headers else 'F6F7F8');tc.append(shading)
                mar=OxmlElement('w:tcMar')
                for edge,value in [('top','65'),('bottom','65'),('left','65'),('right','65')]:
                    node=OxmlElement('w:'+edge);node.set(qn('w:w'),value);node.set(qn('w:type'),'dxa');mar.append(node)
                tc.append(mar)
                for p in cell.paragraphs:
                    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    pf=p.paragraph_format;pf.line_spacing=1.15;pf.space_before=pf.space_after=Pt(0);pf.first_line_indent=Pt(0)
                    pf.keep_with_next=n<len(rows)
                    pp=p._p.get_or_add_pPr();pp.get_or_add_ind().set(qn('w:firstLineChars'),'0');snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pp.append(snap)
                    for run in p.runs:format_run(run,10,values is headers)
        return t

    req=stats['request'];kind=req['report_type'];cycle={'DAY':'日报','WEEK':'周报','MONTH':'月报'}[kind]
    for index,(hid,h) in enumerate(stats['hangars'].items()):
        c,r=h['current'],h['relative'];facts=fact_pack(h)
        title_p=para('空气质量巡测研判'+cycle,24,True,after=8,center=True)
        title_p.paragraph_format.page_break_before=bool(index)
        para((req['start_date'] if kind=='DAY' else req['start_date']+' 至 '+req['end_date'])+'    '+cfg.company,9,after=10,center=True)
        if stats['data_kind']=='synthetic':note('合成数据测试，不用于业务决策。')
        note(h['name']+f"    观测日期 {c['observed_days']}/{c['days']}    有效采样 {c['samples']}条")
        heading('一、空气质量情况')
        points,basis=series(h)
        scope=(f"采样分钟覆盖{points[0][0]}—{points[-1][0]}，仅代表该时段与观测位置。" if basis=='分钟巡测均值' else f"统计结果覆盖{c['observed_days']}个观测日，仅代表已采样的时段与位置。")
        para(facts.get('particles','本期暂无单位明确的有效浓度统计。')+scope,after=7)
        rows=[]
        for m in METRICS:
            v=c['survey'][m]
            vals=[x[m]['value'] for _,x in points if x.get(m,{}).get('value') is not None and x[m].get('unit')==v['unit']]
            span=(survey_value_text(min(vals),v['unit'])+'—'+survey_value_text(max(vals),v['unit'])) if vals else '/'
            rows.append([LABELS[m],v['unit'] or '未确认',survey_value_text(v['value'],v['unit']),span,f"{v['days']}/{c['days']}"])
        table(['监测因子','单位','本期巡测均值',('分钟' if basis=='分钟巡测均值' else '时段' if kind=='DAY' and basis=='时段巡测均值' else '日')+'均值范围','有效日期'],rows,[2.2,2,3.4,6.2,2.8])
        note('注：范围为有观测的'+basis+'范围，非瞬时极值；斜线表示不可计算。气体保留原始体积分数单位。')
        heading('二、采集指标变化趋势')
        data=trend_image(h,runtime)
        if data:
            p=para('',after=3);p.paragraph_format.line_spacing=1;p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(io.BytesIO(data),width=Cm(16.6))
        else:para('本期没有可绘制的有效浓度序列。')
        note('图中仅绘制实际有记录的'+('分钟' if basis=='分钟巡测均值' else '时段' if kind=='DAY' and basis=='时段巡测均值' else '日期')+'；空白与断线表示缺测。不同单位分别绘图，不作数值大小的横向比较。')
        if facts.get('peak'):para(facts['peak'],10,after=0)
        summary_p=para('可比观测分析与工作建议',16,True,after=4)
        summary_p.paragraph_format.page_break_before=True
        note(h['name']+'    '+req['start_date']+' 至 '+req['end_date'])
        para('前后期同条件对比',12,True,before=4,keep=True)
        if r.get('changes'):
            ch=r['changes'];current=r.get('paired_current') or r['current'];previous=r.get('previous')
            para(f"可比观测层{r['usable_strata']}/{r['current_strata']}，覆盖率{r['coverage']:.2%}。下表仅描述这些观测层，不代表全期或区域整体变化。",10,after=6)
            table(['项目','对比期','本期','变化'],[['颗粒物相对指数',f"{ch['previous_value']:.2f}",f"{ch['current_value']:.2f}",f"{ch['net']:+.2f}点"]]+[[LABELS[m]+'指数项',f"{previous['terms'][m]:.2f}" if previous else '/',f"{current['terms'][m]:.2f}",f"{ch['terms'][m]:+.2f}点"] for m in r['panel']], [5.2,3.8,3.8,3.8])
            note(f"对比期：{h['previous']['start']}至{h['previous']['end']}。指数以可比历史基准为100，颗粒物等权；指数项不代表污染源贡献。"+('部分观测层采用短期历史参考。' if any(s['baseline_days']<14 for s in r['strata']) else ''))
        else:
            para('可比历史记录尚不足，本期不列相对指数及其占比。浓度水平与变化特征以上页实测统计为依据。',10)
        heading('三、环境空气质量预测')
        para('尚无与本期对应的正式预报资料，不作未来空气质量等级及扩散条件判断。',10,after=8)
        heading('四、工作建议')
        analysis=analyses[hid]
        paragraphs=analysis.metadata.get('advice',{}).get('paragraphs')
        if analysis.metadata.get('engine')!='ollama':note('模型研判待生成，以下为显式模板模式内容。')
        if paragraphs:
            for n,item in enumerate(paragraphs):
                p=para('',11,before=3,after=8)
                format_run(p.add_run('（'+'一二三'[n]+'）'+item['title']+'。'),11,True)
                format_run(p.add_run(item['paragraph']),11)
        else:
            for s in analysis.sections['advice'].split('\n'):para(s,11,after=8)
        para('资料与统计口径',10,True,before=7,after=4,keep=True)
        notes='日内按有效分钟等权、期间按有数据日期等权统计；缺测不插值。本报告为巡测分析，不作为正式日AQI或区域空气质量评价。'
        if stats['data_kind']=='historical_sensor_replay':notes+='历史记录按项目和传感器归集，任务关联缺失，机库归属未确认；单位依据当前同项目配置解释，历史配置未追溯核验。'
        note(notes)
    target=output/'report.docx';doc.save(target)
    # Preserve opaque source package parts; only document content, relationships,
    # content types and newly generated media may differ from the reference.
    with ZipFile(TEMPLATE) as src,ZipFile(target) as generated:
        parts={n:src.read(n) for n in src.namelist()}
        for n in generated.namelist():
            if n in ('word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml') or n not in parts:parts[n]=generated.read(n)
    with ZipFile(target,'w',ZIP_DEFLATED) as out:
        for n,v in parts.items():out.writestr(n,v)
    loaded=Document(target)
    if len(loaded.tables)!=len(doc.tables):raise ValueError('Report table readback failed')
    return {'template_version':VERSION,'template_sha256':sha256(TEMPLATE.read_bytes()).hexdigest(),'readback':'passed','report_type':kind,'runtime':{k:v for k,v in runtime.items() if k!='env'}}
