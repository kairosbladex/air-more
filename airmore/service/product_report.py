"""Render database-backed reports with the approved product-format layout.

The service report contains device observations, not a regional bulletin.  The
Word structure and visual language nevertheless follow the user-approved
product template.  Unsupported formal AQI, forecast, ranking, and source
claims remain explicit gaps in the rendered report.
"""
from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path
import re

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from airmore.regional.product_style import (
    TEMPLATE,
    TEMPLATE_SHA256,
    TEMPLATE_VERSION,
    heading,
    image,
    new_document,
    page_break,
    paragraph,
    table,
)

from .models import LABELS, METRICS

VERSION = TEMPLATE_VERSION + "+service/v2"
COLORS = ('#5B9BD5', '#ED7D31', '#70AD47', '#4472C4', '#A5A5A5', '#FFC000')


def survey_value_text(value, unit):
    # Keep the renderer deployable on the stable service without changing
    # request models, database reads, or device-linkage validation.
    return format(value, '.3f' if unit in ('ppm', 'ppb') else '.2f') if value is not None else '/'


def _date_text(value: str) -> str:
    current = date.fromisoformat(value)
    return f'{current.month}月{current.day}日'


def _date_range(start: str, end: str) -> str:
    a, b = date.fromisoformat(start), date.fromisoformat(end)
    if a == b:
        return f'{a.year}年{a.month:02d}月{a.day:02d}日'
    return f'{a.year}年{a.month:02d}月{a.day:02d}日—{b.month:02d}月{b.day:02d}日'


def _current_days(hangar: dict):
    current = hangar['current']
    return [(day, value) for day, value in hangar['days'].items()
            if current['start'] <= day <= current['end']]


def _formal_aqi_rows(hangars: dict):
    rows = []
    for hangar in hangars.values():
        for day, value in _current_days(hangar):
            standard = value.get('standard', {})
            if standard.get('aqi') is not None and not standard.get('partial'):
                rows.append((hangar['name'], day, standard))
    return rows


def _quality(aqi):
    if aqi is None:
        return '不可计算'
    return next((name for limit, name in (
        (50, '优'), (100, '良'), (150, '轻度污染'),
        (200, '中度污染'), (300, '重度污染'), (500, '严重污染')) if aqi <= limit), '不可计算')


def _pollutants(standard):
    names = [LABELS[m] for m in standard.get('primary_pollutants', [])]
    return '、'.join(names) if names else '/'


def _write_chart_font(runtime):
    return [runtime.get('font_family') or runtime.get('docx_font_family') or 'DejaVu Sans',
            'DejaVu Sans']


def _save_chart(fig, path: Path):
    fig.savefig(path, dpi=190, bbox_inches='tight', facecolor=fig.get_facecolor())
    plt.close(fig)


def _make_share_chart(hangars: dict, path: Path, runtime: dict):
    totals = {metric: 0 for metric in METRICS}
    eligible_days = 0
    formal_rows = _formal_aqi_rows(hangars)
    for _, _, standard in formal_rows:
        eligible_days += 1
        for metric, value in standard.get('iaqi', {}).items():
            if metric in totals:
                totals[metric] += value

    with plt.rc_context({'font.family': _write_chart_font(runtime)}):
        # The reference report uses two compact horizontal panels.  Keeping
        # that aspect ratio prevents the placeholder panel from consuming a
        # full page when formal AQI data is unavailable.
        fig, axes = plt.subplots(1, 2, figsize=(8, 3.25))
        values = [totals[metric] for metric in METRICS]
        latest = formal_rows[-1][2] if formal_rows else None
        panels = [
            (values, f'本期有效日分指数累计占比（{eligible_days}个设备日）'),
            ([latest.get('iaqi', {}).get(metric, 0) for metric in METRICS] if latest else [],
             f'{formal_rows[-1][1]}分指数构成' if latest else '完整日评价资料不足'),
        ]
        legend_wedges = None
        for ax, (panel_values, title) in zip(axes, panels):
            if sum(panel_values):
                wedges, _, _ = ax.pie(panel_values, colors=COLORS, startangle=90,
                                      autopct='%1.1f%%', pctdistance=.76,
                                      textprops={'fontsize': 8})
                legend_wedges = wedges
            else:
                ax.pie([1], colors=['#EEEEEE'],
                       wedgeprops={'edgecolor': '#BFBFBF', 'linewidth': .7})
                ax.text(0, 0, '完整日评价资料不足\n暂不计算分指数占比',
                        ha='center', va='center', fontsize=9, color='#606060')
            ax.set_title(title, fontsize=10)
            ax.set_aspect('equal')
        if legend_wedges is None:
            legend_wedges = [plt.Rectangle((0, 0), 1, 1, color=color) for color in COLORS]
        fig.legend(legend_wedges, [LABELS[m] for m in METRICS], loc='lower center',
                   ncol=6, frameon=False, fontsize=8)
        fig.subplots_adjust(bottom=.18, wspace=.18)
        _save_chart(fig, path)


def _make_aqi_chart(hangars: dict, path: Path, runtime: dict):
    all_days = sorted({day for hangar in hangars.values() for day, _ in _current_days(hangar)})
    x = list(range(len(all_days)))
    with plt.rc_context({'font.family': _write_chart_font(runtime), 'axes.unicode_minus': False}):
        fig, ax = plt.subplots(figsize=(8, 2.6))
        plotted = 0
        for hangar in hangars.values():
            by_day = {day: value.get('standard', {}).get('aqi') for day, value in _current_days(hangar)}
            values = [by_day.get(day) if by_day.get(day) is not None else float('nan') for day in all_days]
            if any(value == value for value in values):
                ax.plot(x, values, marker='o', markersize=3.5, linewidth=1.5, label=hangar['name'])
                plotted += 1
        if plotted:
            ax.axhline(100, color='#ED7D31', linestyle='--', linewidth=1)
            step = max(1, len(all_days) // 8)
            ticks = sorted(set([0, len(all_days) - 1, *range(0, len(all_days), step)]))
            ax.set_xticks(ticks, [all_days[i][5:].replace('-', '/') for i in ticks], fontsize=8)
            ax.set_ylabel('AQI')
            ax.set_ylim(bottom=0)
            ax.grid(axis='y', alpha=.2)
            if len(hangars) <= 8:
                ax.legend(loc='upper left', bbox_to_anchor=(0, 1.02), ncol=2,
                          frameon=False, fontsize=8)
        else:
            ax.text(.5, .5, '没有满足完整日评价条件的AQI序列',
                    ha='center', va='center', transform=ax.transAxes, color='#606060')
            ax.set_axis_off()
        _save_chart(fig, path)


def _metric_points(hangar: dict):
    current = hangar['current']
    if hangar.get('report_type') == 'DAY':
        # A single daily report still gets a stable one-point chart.  The
        # minute/hour series stays available in the evidence package and is
        # not silently presented as a full-day regulatory average.
        return [(current['end'], current['survey'])]
    return [(day, value['survey']) for day, value in _current_days(hangar)]


def _make_trend_chart(hangars: dict, path: Path, runtime: dict):
    all_days = sorted({day for hangar in hangars.values() for day, _ in _metric_points(hangar)})
    x = list(range(len(all_days)))
    with plt.rc_context({'font.family': _write_chart_font(runtime), 'axes.unicode_minus': False}):
        fig, axes = plt.subplots(3, 2, figsize=(8, 8.2), facecolor='#17212F')
        for index, (ax, metric) in enumerate(zip(axes.flat, METRICS)):
            ax.set_facecolor('#17212F')
            plotted = 0
            for hangar in hangars.values():
                by_day = {day: values.get(metric, {}) for day, values in _metric_points(hangar)}
                units = {v.get('unit') for v in by_day.values() if v.get('value') is not None}
                if len(units) != 1:
                    continue
                values = [by_day.get(day, {}).get('value') for day in all_days]
                if not any(value is not None for value in values):
                    continue
                ax.plot(x, [value if value is not None else float('nan') for value in values],
                        color=COLORS[index], marker='o', markersize=3, linewidth=1.5,
                        label=hangar['name'])
                plotted += 1
            ax.set_title(LABELS[metric] + ' 巡测统计',
                         color='white', fontsize=11)
            unit = next((v.get('unit') for hangar in hangars.values()
                         for _, values in _metric_points(hangar)
                         for v in [values.get(metric, {})] if v.get('unit')), '/')
            ax.set_ylabel(unit, color='#CBD5E1', fontsize=9)
            if all_days:
                step = max(1, len(all_days) // 6)
                ticks = sorted(set([0, len(all_days) - 1, *range(0, len(all_days), step)]))
                ax.set_xticks(ticks, [all_days[i][5:].replace('-', '/') for i in ticks], fontsize=8)
            ax.tick_params(colors='#CBD5E1')
            ax.grid(alpha=.15)
            for spine in ax.spines.values():
                spine.set_color('#607080')
            if plotted == 0:
                ax.text(.5, .5, '无可用数据', ha='center', va='center',
                        transform=ax.transAxes, color='#CBD5E1')
            elif len(hangars) <= 8:
                ax.legend(loc='upper left', fontsize=7, frameon=False,
                          labelcolor='white')
        fig.tight_layout(pad=2)
        _save_chart(fig, path)


def _advice_paragraphs(analyses):
    result = []
    seen = set()
    for analysis in analyses.values():
        metadata = getattr(analysis, 'metadata', {}) or {}
        items = metadata.get('advice', {}).get('paragraphs', [])
        candidates = [item.get('paragraph', '') for item in items if item.get('paragraph')]
        if not candidates:
            candidates = str(getattr(analysis, 'sections', {}).get('advice', '')).splitlines()
        for candidate in candidates:
            candidate = candidate.strip()
            if candidate and candidate not in seen:
                seen.add(candidate)
                result.append(candidate)
    return result


def _overview(stats, analyses):
    parts = []
    for hid, hangar in stats['hangars'].items():
        analysis = analyses.get(hid)
        text = getattr(analysis, 'sections', {}).get('overview', '') if analysis else ''
        if text:
            parts.append(f"{hangar['name']}：{text.replace(chr(10), ' ')}")
    parts.append('本报告仅代表实际采样覆盖，不替代固定点完整日评价或区域空气质量评价。')
    if stats.get('data_kind') == 'historical_sensor_replay':
        parts.append('历史记录按项目和传感器归集，当前任务关联未能确认历史机库归属。')
    return ' '.join(parts)


def _unit_note(stats):
    units = set()
    for hangar in stats['hangars'].values():
        for value in hangar['current']['survey'].values():
            if value.get('unit'):
                units.add(value['unit'])
    if not units:
        return '单位或有效值不足，斜线表示不可计算，不表示零。'
    return '各指标保留已核验的原生单位；缺测、单位冲突或无效值以斜线表示，不补零。'


def render_document(stats, analyses, composite, cfg, output, runtime):
    """Render one service report using the shared product-format template."""
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    assets = output / 'assets'
    assets.mkdir(exist_ok=True)

    request = stats['request']
    start, end = request['start_date'], request['end_date']
    hangars = stats['hangars']
    title_period = _date_range(start, end)
    footer = f'{cfg.company}空气质量研判报告'
    doc = new_document('空气质量', footer_title=footer,
                       font_name=runtime.get('docx_font_family'))
    center = WD_ALIGN_PARAGRAPH.CENTER

    for _ in range(5):
        paragraph(doc, '', space=12)
    paragraph(doc, '空气质量研判报告', 36, align=center, space=18)
    paragraph(doc, title_period, 14, align=center, space=12)
    paragraph(doc, '报告类型：' + {'DAY': '日报', 'WEEK': '周报', 'MONTH': '月报'}[request['report_type']], 10.5, align=center)
    page_break(doc)

    heading(doc, '空气质量情况')
    paragraph(doc, _overview(stats, analyses))
    heading(doc, '（一）空气质量情况表', True)
    value_rows = []
    for hangar in hangars.values():
        current = hangar['current']
        values = [survey_value_text(current['survey'][metric]['value'], current['survey'][metric]['unit'])
                  + ('\n' + current['survey'][metric]['unit']
                     if current['survey'][metric]['value'] is not None and current['survey'][metric]['unit'] else '')
                  for metric in METRICS]
        value_rows.append([hangar['name'], *values, current['samples'],
                           f"{current['observed_days']}/{current['days']}"])
    values_table = table(doc, ['采集设备'] + [LABELS[m] for m in METRICS] + ['有效样本', '有效日期'],
          value_rows or [['/', *(['/'] * len(METRICS)), '/', '/']],
          [2.5, 1.5, 1.5, 1.4, 1.7, 1.5, 1.5, 1.35, 1.7], 8.5)
    # Make room for decimal values even when Word substitutes a wider font.
    margins = OxmlElement('w:tblCellMar')
    for edge in ('left', 'right'):
        margin = OxmlElement('w:' + edge)
        margin.set(qn('w:w'), '40')
        margin.set(qn('w:type'), 'dxa')
        margins.append(margin)
    values_table._tbl.tblPr.append(margins)
    paragraph(doc, _unit_note(stats), 8.5)
    coverage_rows = []
    for hangar in hangars.values():
        current, previous = hangar['current'], hangar['previous']
        coverage_rows.append([hangar['name'], f"{current['observed_days']}/{current['days']}",
                              str(current['samples']), f"{previous['observed_days']}/{previous['days']}",
                              str(previous['samples'])])
    heading(doc, '采样覆盖与质量状态', True)
    table(doc, ['采集设备', '本期有数据日', '本期样本', '比较期有数据日', '比较期样本'],
          coverage_rows or [['/', '/', '/', '/', '/']], [3.4, 2.5, 2.4, 2.8, 2.4], 9)
    paragraph(doc, '没有记录的日期不补零；正式AQI、综合指数、跨设备排名和污染源贡献率仅在满足相应业务口径时计算。', 8.5)

    page_break(doc)
    heading(doc, '（二）六项污染物分指数占比分析', True)
    if _formal_aqi_rows(hangars):
        paragraph(doc, '仅汇总满足完整日评价条件的设备日分指数，用于识别本期指标关注重点；该占比不是污染源贡献率。')
    else:
        paragraph(doc, '当前没有满足完整日评价条件的设备日，暂不计算六项污染物分指数占比；巡测相对指数也不代表正式综合指数。')
    share_path = assets / 'service-iaqi-shares.png'
    _make_share_chart(hangars, share_path, runtime)
    image(doc, share_path)
    aqi_rows = []
    for name, day, standard in _formal_aqi_rows(hangars)[-8:]:
        aqi_rows.append([name, _date_text(day), standard['aqi'], _quality(standard['aqi']), _pollutants(standard)])
    table(doc, ['采集设备', '日期', 'AQI', '空气质量', '首要污染物'],
          aqi_rows or [['/', '/', '/', '不可计算', '/']], [3.5, 2.3, 1.7, 3.2, 3.0], 9)
    paragraph(doc, '正式AQI仅列满足完整日评价条件的记录；其余日期保留在证据包中，不以缺测推断正常或污染。', 8.5)
    heading(doc, '近期空气质量变化', True)
    aqi_path = assets / 'service-aqi-trend.png'
    _make_aqi_chart(hangars, aqi_path, runtime)
    image(doc, aqi_path)

    page_break(doc)
    heading(doc, '采集指标变化趋势分析')
    paragraph(doc, f'{title_period}六项污染物巡测统计变化如下。气体与颗粒物的单位不同，各图独立纵轴，不作数值大小的横向比较。')
    trend_path = assets / 'service-trends.png'
    _make_trend_chart(hangars, trend_path, runtime)
    image(doc, trend_path)
    paragraph(doc, '图中空白或断线表示没有可用记录；单日巡测结果只代表实际采样时段和位置。', 8.5)

    page_break(doc)
    heading(doc, '环境空气质量预测')
    heading(doc, '（一）空气质量形势预报', True)
    end_date = date.fromisoformat(end)
    for offset in range(3):
        future = end_date + timedelta(days=offset + 1)
        paragraph(doc, f'{future.month}月{future.day}日，正式预报数据待接入；本报告不作未来空气质量等级、AQI数值或扩散条件判断。')
    heading(doc, '（二）未来五天AQI数值预报', True)
    forecast_dates = [end_date + timedelta(days=offset + 1) for offset in range(5)]
    table(doc, [f'{d.month}月{d.day}日' for d in forecast_dates],
          [['/' for _ in forecast_dates], ['/' for _ in forecast_dates],
           ['/' for _ in forecast_dates], ['/' for _ in forecast_dates]],
          [2.92] * 5, 9)
    paragraph(doc, '服务当前未接入与本期对应的正式预报资料，表中斜线表示不可计算，不表示空气质量为零或已判定为优良。', 8.5)

    page_break(doc)
    heading(doc, '工作建议')
    advice = _advice_paragraphs(analyses)
    if not advice:
        advice = ['先核对本期采样覆盖、指标单位和设备关联，再安排同条件复测与结果复核；现有资料不足以判定污染来源。']
    for text in advice:
        p = paragraph(doc, text, 10.5, space=13)
        p.paragraph_format.first_line_indent = Pt(21)
    paragraph(doc, '资料说明：本报告由业务采样记录生成，缺测不插值；结果仅代表实际采样覆盖。报告中的正式AQI、预测、综合指数、排名和污染源归因均以独立业务口径和完整证据为前提。', 8.5)

    visible = '\n'.join(t.text for t in doc._element.body.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
    if re.search(r'https?://|\[\d+\]|公开|数据来源|来源索引|外部网站|GPU|qwen|DataV', visible):
        raise ValueError('source annotation leaked into report')

    target = output / 'report.docx'
    doc.save(target)
    reread = Document(target)
    if not reread.paragraphs or not reread.tables:
        raise ValueError('Product-format report readback failed')
    return {
        'template_version': VERSION,
        'template_sha256': TEMPLATE_SHA256,
        'readback': 'passed',
        'report_type': request['report_type'],
        'runtime': {key: value for key, value in runtime.items() if key != 'env'},
        'layout': 'user-approved-product-format',
        'template_path': str(TEMPLATE),
    }
