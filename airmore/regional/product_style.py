"""Shared typography and tables for the approved air-quality product report."""
import hashlib
from pathlib import Path
import sys
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# This is the user-approved product-format reference.  Both the regional
# renderer and the database-backed service import this same package asset so
# test and production cannot silently drift to different Word layouts.
TEMPLATE = Path(__file__).resolve().parents[1] / 'service/assets/product-report-template.docx'
TEMPLATE_VERSION = 'product-air-quality/2'
TEMPLATE_SHA256 = hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()
DEFAULT_FONT = '宋体-简' if sys.platform == 'darwin' else 'Noto Serif CJK SC'

def font(run,size=10.5,bold=False,name=None):
 name = name or DEFAULT_FONT
 run.font.name=name;run.font.size=Pt(size);run.font.bold=bold
 r_pr = run._element.get_or_add_rPr()
 r_fonts = r_pr.rFonts
 for attribute in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
  r_fonts.set(qn('w:'+attribute),name)

def _font_name(doc, name=None):
 return name or getattr(doc, '_product_font_name', DEFAULT_FONT)

def paragraph(doc,text='',size=10.5,bold=False,align=None,space=6,font_name=None):
 p=doc.add_paragraph();p.paragraph_format.line_spacing=Pt(max(14,size*1.4));p.paragraph_format.space_after=Pt(space)
 if getattr(doc,'_pending_page_break',False):
  p.paragraph_format.page_break_before=True;doc._pending_page_break=False
 snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(snap)
 p.paragraph_format.widow_control=True
 if align is not None:p.alignment=align
 font(p.add_run(text),size,bold,_font_name(doc,font_name))
 return p

def page_break(doc):
 doc._pending_page_break=True

def heading(doc,text,sub=False,font_name=None):
 p=paragraph(doc,text,12 if sub else 16,True,space=8,font_name=font_name)
 p.paragraph_format.keep_with_next=True
 return p

def table(doc,headers,rows,widths,size=9,font_name=None):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for data,row in [(headers,t.rows[0])]+[(d,t.add_row()) for d in rows]:
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  for c,txt,w in zip(row.cells,data,widths):
   c.width=Cm(w);p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.05
   font(p.add_run(str(txt)),size,data is headers,_font_name(doc,font_name))
   if data is headers:
    s=OxmlElement('w:shd');s.set(qn('w:fill'),'E7E6E6');c._tc.get_or_add_tcPr().append(s)
  if data is headers:pr.append(OxmlElement('w:tblHeader'))
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'C8C8C8');borders.append(el)
 t._tbl.tblPr.append(borders)
 return t

def image(doc,path,width=14.5):
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(4);p.add_run().add_picture(str(path),width=Cm(width))


def new_document(region, template=TEMPLATE, footer_title=None, font_name=None):
 d=Document(template)
 d._product_font_name = _font_name(d,font_name)
 normal=d.styles['Normal']._element.find(qn('w:pPr'))
 if normal is not None:
  for element in list(normal):normal.remove(element)
 normal_style=d.styles['Normal']
 normal_style.font.name=d._product_font_name
 normal_style._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),d._product_font_name)
 normal_style._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),d._product_font_name)
 normal_style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),d._product_font_name)
 normal_style._element.get_or_add_rPr().rFonts.set(qn('w:cs'),d._product_font_name)
 for element in list(d._element.body):
  if element.tag!=qn('w:sectPr'):d._element.body.remove(element)
 for section in d.sections:
  grid=section._sectPr.find(qn('w:docGrid'))
  if grid is not None:section._sectPr.remove(grid)
  section.page_width=Cm(21);section.page_height=Cm(29.7)
  section.top_margin=Cm(2.1);section.bottom_margin=Cm(2)
  section.left_margin=Cm(3.175);section.right_margin=Cm(3.175)
  for p in section.footer.paragraphs:p.clear()
  p=section.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  font(p.add_run(footer_title or (region+'空气质量研判报告')+'  ·  '),8,name=d._product_font_name)
  field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
 return d
