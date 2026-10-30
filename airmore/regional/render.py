"""Regional typography, verified DOCX output, and isolated PDF conversion."""

from __future__ import annotations

import hashlib
import io
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from zipfile import ZipFile

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.ft2font import FT2Font
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from airmore.render.docx_render import new_document, _set_eastasia
from .bundle import Bundle, FIELDS, LABELS, require


def preflight(soffice=None) -> dict:
    requested = str(soffice) if soffice else None
    if requested:
        executable = requested if Path(requested).is_file() else shutil.which(requested)
    else:
        executable = shutil.which("soffice")
        if not executable and Path("/opt/homebrew/bin/soffice").is_file():
            executable = "/opt/homebrew/bin/soffice"
    require(bool(executable), "LibreOffice soffice unavailable; install it before generating PDF")
    executable = str(Path(executable).resolve())
    selected = None
    for family in ("Songti SC", "SimSun", "Noto Serif CJK SC", "PingFang SC", "Noto Sans CJK SC"):
        try:
            p = font_manager.findfont(font_manager.FontProperties(family=[family]), fallback_to_default=False)
            face = FT2Font(p)
            if all(face.get_char_index(ord(c)) for c in "环境空气质量研判日报μ³℃"):
                selected = (family, p)
                break
        except (ValueError, RuntimeError):
            continue
    require(selected is not None, "Chinese font unavailable; install Songti SC, SimSun or Noto CJK")
    env = os.environ.copy()
    if not env.get("FONTCONFIG_FILE") and Path("/opt/homebrew/etc/fonts/fonts.conf").is_file():
        env["FONTCONFIG_FILE"] = "/opt/homebrew/etc/fonts/fonts.conf"
    version = subprocess.run([executable, "--version"], check=True, capture_output=True,
                             text=True, timeout=20, env=env).stdout.strip()
    # Native macOS LibreOffice resolves these CJK faces by their localized names.
    docx_family = ({"Songti SC": "宋体-简", "PingFang SC": "苹方-简"}.get(selected[0], selected[0])
                   if sys.platform == "darwin" else selected[0])
    return {"soffice": executable, "version": version, "font_family": selected[0],
            "docx_font_family": docx_family, "font_path": str(selected[1]), "env": env}


def chart_bytes(chart: dict, runtime: dict) -> bytes:
    font = font_manager.FontProperties(fname=runtime["font_path"])
    colors = ["#5B9BD5", "#ED7D31", "#70AD47", "#4472C4", "#A5A5A5", "#FFC000"]
    with plt.rc_context({"font.family": runtime["font_family"], "axes.unicode_minus": False}):
        fig, axes = plt.subplots(1, 2, figsize=(8, 3.25), dpi=180)
        try:
            for ax, panel in zip(axes, chart["panels"]):
                ax.set_title("昨日" if panel["period"] == "day" else "当月累计", fontproperties=font, fontsize=12)
                if panel["ratios"]:
                    values = [panel["ratios"][f] for f in FIELDS]
                    wedges, _, _ = ax.pie(values, colors=colors, autopct=lambda p: f"{p:.1f}%" if p else "",
                                          textprops={"fontsize": 9}, startangle=90)
                    ax.legend(wedges, [LABELS[f] for f in FIELDS], loc="lower center",
                              bbox_to_anchor=(0.5, -0.2), ncol=3, fontsize=8, frameon=False)
                else:
                    reason = panel["reason"]
                    reason = "\n".join(reason[i:i + 16] for i in range(0, len(reason), 16))
                    ax.text(.5, .5, reason, ha="center", va="center", fontproperties=font, fontsize=11)
                    ax.set_axis_off()
                ax.set_aspect("equal")
            fig.subplots_adjust(left=.025, right=.975, top=.89, bottom=.16, wspace=.18)
            out = io.BytesIO()
            fig.savefig(out, format="png", dpi=180)
            return out.getvalue()
        finally:
            plt.close(fig)


def _shade(cell, color):
    el = OxmlElement("w:shd")
    el.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(el)


def _table(doc, block, family):
    table = doc.add_table(rows=1, cols=len(block["header"]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = block["widths"]
    require(sum(widths) <= 14.65 + .001, "table exceeds reference text width")
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "bottom", "left", "right", "insideH", "insideV"):
        edge = OxmlElement("w:" + side)
        for k, v in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
            edge.set(qn("w:" + k), v)
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    for i, data in enumerate([block["header"], *block["rows"]]):
        row = table.rows[0] if i == 0 else table.add_row()
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        style = block["styles"][i - 1] if i > 0 and block["styles"] else {}
        for j, (cell, content) in enumerate(zip(row.cells, data)):
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = content
            margins = OxmlElement("w:tcMar")
            for side in ("top", "bottom", "left", "right"):
                margin = OxmlElement("w:" + side)
                margin.set(qn("w:w"), "45")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            cell._tc.get_or_add_tcPr().append(margins)
            if i == 0 or style.get("fill"):
                _shade(cell, "9BAC84" if i == 0 else style["fill"])
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    _set_eastasia(run, "Times New Roman", family)
                    run.font.size = Pt(9)
                    run.font.bold = i == 0
                    run.font.color.rgb = RGBColor.from_string(style.get("color") or "000000")


def render_docx(b: Bundle, view: dict, target: Path, runtime: dict) -> dict[str, bytes]:
    doc = new_document()
    family = runtime["docx_font_family"]
    normal = doc.styles["Normal"]
    _set_eastasia(normal, "Times New Roman", family)
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(5)
    for name, size in (("Heading 1", 15), ("Heading 2", 13)):
        st = doc.styles[name]
        _set_eastasia(st, "Times New Roman", family)
        fonts = st.element.rPr.rFonts
        for attr in list(fonts.attrib):
            if attr.lower().endswith("theme"):
                del fonts.attrib[attr]
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.keep_with_next = True
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(3.175)
    footer = sec.footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run(view["disclaimer"] + "｜第 ")
    r.font.size = Pt(8)
    _set_eastasia(r, "Times New Roman", family)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页").font.size = Pt(8)
    charts = {c["id"]: chart_bytes(c, runtime) for c in view["charts"]}
    images = {r["id"]: b.assets[r["path"]] for r in b.images}
    media = charts | images
    for block in view["blocks"]:
        kind = block["type"]
        if kind == "table":
            _table(doc, block, family)
        elif kind in {"chart", "weather"}:
            ids = [block["id"]] if kind == "chart" else block["ids"]
            if kind == "weather":
                table = doc.add_table(rows=1, cols=2)
                table.autofit = False
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                for col in table.columns:
                    col.width = Cm(7.2)
                paragraphs = [c.paragraphs[0] for c in table.rows[0].cells]
            else:
                paragraphs = [doc.add_paragraph()]
            for iid, p in zip(ids, paragraphs):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                width = 14.4 if kind == "chart" else 6.8
                with Image.open(io.BytesIO(media[iid])) as img:
                    width = min(width, 8.0 * img.width / img.height)
                p.add_run().add_picture(io.BytesIO(media[iid]), width=Cm(width))
                p.paragraph_format.keep_with_next = True
            if kind == "weather":
                for cell, caption in zip(table.rows[0].cells, block["captions"]):
                    p = cell.add_paragraph(caption)
                    p.paragraph_format.keep_together = True
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.1
                    for r in p.runs:
                        _set_eastasia(r, "Times New Roman", family)
                        r.font.size = Pt(8.5)
                        r.font.color.rgb = RGBColor(80, 80, 80)
        else:
            style = {"section": "Heading 1", "heading": "Heading 2"}.get(kind)
            p = doc.add_paragraph(block["text"], style=style)
            p.paragraph_format.keep_together = True
            if block.get("page_break_before"):
                p.paragraph_format.page_break_before = True
            if kind in {"title", "issue", "caption"}:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if kind == "date":
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                border = OxmlElement("w:pBdr")
                bottom = OxmlElement("w:bottom")
                for k, v in {"val": "single", "sz": "18", "color": "FF0000"}.items():
                    bottom.set(qn("w:" + k), v)
                border.append(bottom)
                p._p.get_or_add_pPr().append(border)
            if kind in {"title", "issue", "date", "caption"}:
                p.paragraph_format.keep_with_next = True
            if kind == "note":
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.1
            for r in p.runs:
                if kind == "title":
                    r.font.size = Pt(34)
                    r.font.color.rgb = RGBColor(255, 0, 0)
                elif kind in {"issue", "date"}:
                    r.font.size = Pt(15)
                elif kind == "caption":
                    r.font.size = Pt(11)
                    r.font.bold = True
                elif kind == "note":
                    r.font.size = Pt(8.5)
                    r.font.color.rgb = RGBColor(80, 80, 80)
    doc.save(target)
    return media


def verify_docx(path: Path, view: dict, media: dict[str, bytes]) -> None:
    doc = Document(path)
    expected = []
    for block in view["blocks"]:
        if block["type"] == "table":
            expected.append(("table", [block["header"], *block["rows"]]))
        elif block["type"] in {"chart", "weather"}:
            ids = [block["id"]] if block["type"] == "chart" else block["ids"]
            captions = block.get("captions", [])
            expected.append(("media", [hashlib.sha256(media[k]).hexdigest() for k in ids], captions))
        elif "text" in block:
            expected.append(("text", block["text"]))
    actual = []
    for el in doc.iter_inner_content():
        blips = list(el._element.iter(qn("a:blip")))
        if blips:
            if hasattr(el, "rows"):
                captions = [c.text.removeprefix("\n") for row in el.rows for c in row.cells if c.text]
            else:
                require(not el.text, "unexpected text in media block")
                captions = []
            actual.append(("media", [hashlib.sha256(doc.part.related_parts[
                node.get(qn("r:embed"))].blob).hexdigest() for node in blips], captions))
        elif hasattr(el, "rows"):
            actual.append(("table", [[c.text for c in row.cells] for row in el.rows]))
        elif el.text:
            actual.append(("text", el.text))
    require(actual == expected, "DOCX readback differs from bound report facts")
    require(view["disclaimer"] in doc.sections[0].footer.paragraphs[0].text, "DOCX disclaimer missing")
    expected_hashes = {hashlib.sha256(v).hexdigest() for v in media.values()}
    with ZipFile(path) as z:
        actual_hashes = {hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith("word/media/")}
    require(actual_hashes == expected_hashes, "DOCX chart or weather image differs from bound asset")


def convert_pdf(docx: Path, runtime: dict) -> Path:
    pdf = docx.with_suffix(".pdf")
    require(not pdf.exists(), "PDF target must be new")
    with tempfile.TemporaryDirectory(prefix="airmore-lo-") as profile:
        subprocess.run([runtime["soffice"], f"-env:UserInstallation={Path(profile).as_uri()}",
                        "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
                       check=True, capture_output=True, timeout=120, env=runtime["env"])
    require(pdf.is_file() and pdf.stat().st_size > 100, "PDF conversion produced no new output")
    require(pdf.read_bytes().startswith(b"%PDF-"), "invalid PDF output")
    return pdf
