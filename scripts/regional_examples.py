"""Create reproducible synthetic, missing-data, and reference-transcription bundles."""

from __future__ import annotations

import argparse
import csv
import json
import sys
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from airmore.regional.bundle import (CONTRIBUTION_COLUMNS, FIELDS, RANK_COLUMNS,
                                     STAT_COLUMNS, UNITS)


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def write_csv(path, rows, columns):
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=sorted(columns))
        writer.writeheader()
        writer.writerows(rows)


def make_example(root: Path, count=4, missing=False, report_date="2026-09-01", reference: Path | None = None):
    root.mkdir(parents=True, exist_ok=True)
    d = date.fromisoformat(report_date)
    yesterday = d - timedelta(days=1)
    kind = "reference_transcription" if reference else "synthetic"
    manifest = {
        "schema_version": "1.0", "data_kind": kind,
        "region": {"id": "demo", "name": "东昌府区参考件回放" if reference else "示例区域"},
        "report_date": report_date, "issue": "244" if reference else "TEST001",
        "published_at": report_date + "T09:00:00+08:00", "table_scope_id": "district",
        "sources": [{"id": "S1", "kind": kind,
                     "label": "第244期参考件转录（时刻为回放约定，非核实的原始发布时间）" if reference else "确定性合成测试数据",
                     "issued_at": report_date + "T08:00:00+08:00"}],
        "stations": [], "files": {k: k + (".json" if k in {"forecasts", "advice"} else ".csv")
                                   for k in ("statistics", "ranks", "contributions", "forecasts", "advice")},
        "weather_images": [],
    }
    reference_tables = None
    if reference:
        from docx import Document
        doc = Document(reference)
        reference_tables = {
            "day": [[c.text.strip() for c in r.cells] for r in doc.tables[0].rows[1:]],
            "month": [[c.text.strip() for c in r.cells] for r in doc.tables[2].rows[1:]],
        }
        names = [r[0] for r in reference_tables["day"]]
        count = len(names)
    else:
        names = [f"示例站点{i + 1}" for i in range(count)]
        if count > 1:
            names[-1] = "示例区域东部产业园空气质量综合监测固定站"
    ids = {}
    for i, name in enumerate(names):
        eid = f"P{i + 1:02d}"
        ids[name] = eid
        focus = name in {"凤凰工业园", "嘉明经济开发区"} if reference else i < min(2, count)
        manifest["stations"].append({"id": eid, "name": name, "region_id": "demo", "order": i,
                                     "focus": focus, "kind": "park" if focus else "station"})
    stats, ranks, contributions, forecasts, advice = [], [], [], [], []

    def base(eid, period, metric, suffix):
        return {"id": f"{eid}_{period}_{metric}_{suffix}", "entity_id": eid, "period": period,
                "start": (yesterday if period == "day" else yesterday.replace(day=1)).isoformat(),
                "end": yesterday.isoformat(), "metric": metric, "quality": "valid", "reason": "",
                "source_id": "S1", "basis": "参考件原值转录，不反推正式算法" if reference else "合成核定结果，仅测试展示"}

    for period in ("day", "month"):
        rows = reference_tables[period] if reference else [None] * count
        for i, original in enumerate(rows):
            eid = ids[original[0]] if original else f"P{i + 1:02d}"
            for j, metric in enumerate((*FIELDS, "composite")):
                r = base(eid, period, metric, "stat")
                r["unit"] = UNITS[metric]
                r["value"] = original[j + 1] if original else ([8 + i, 17 + i, .4 + i / 10, 144 + i, 28 + i, 7 + i, 2.16 + i / 10][j])
                if r["value"] == "/" or (missing and (period == "month" or i == 0)):
                    r.update(value="", quality="missing", reason="参考件标注无数据" if reference else "合成缺测场景")
                stats.append(r)
            for metric in (("composite",) if reference else (*FIELDS, "composite")):
                r = base(eid, period, metric, "rank")
                r.update(rank=original[8] if original else i + 1, total="" if reference else count,
                         scope_id="district", scope_name="主城四区（参考件口径）" if reference else "示例区域全部参评对象")
                if missing:
                    r.update(rank="", quality="missing", reason="未提供核定排名")
                ranks.append(r)
            station = next(s for s in manifest["stations"] if s["id"] == eid)
            if station["focus"]:
                for j, metric in enumerate(FIELDS):
                    r = base(eid, period, metric, "contribution")
                    r.update(mode="ratios", value=[.05, .16, .06, .36, .18, .19][j])
                    if reference:
                        # The reference bitmap labels are transcribed as displayed, not reverse-calculated.
                        if station["name"] == "嘉明经济开发区":
                            r["value"] = [.05, .16, .05, .36, .19, .19][j]
                        if period == "day":
                            r.update(value="", quality="missing", reason="参考件注明省平台无数据")
                    elif missing:
                        r.update(value="", quality="missing", reason="六项贡献率未提供")
                    contributions.append(r)
    if reference:
        for s in manifest["stations"]:
            if s["focus"]:
                for f in FIELDS:
                    r = base(s["id"], "day", f, "rank")
                    r.update(rank="", total="", scope_id="district", scope_name="主城四区（参考件口径）",
                             quality="missing", reason="参考件单项排名为空")
                    ranks.append(r)
                r = base(s["id"], "month", "composite", "city_rank")
                r.update(rank=125 if s["name"] == "凤凰工业园" else 85, total="",
                         scope_id="city", scope_name="全市（参考件口径）")
                ranks.append(r)
        combined = {"id": "COMBINED", "name": "两地合并", "region_id": "demo", "order": 99,
                    "focus": False, "kind": "combined"}
        manifest["stations"].append(combined)
        r = base("COMBINED", "month", "composite", "stat")
        r.update(value=2.53, unit="1")
        stats.append(r)
        r = base("COMBINED", "month", "composite", "city_rank")
        r.update(rank=112, total="", scope_id="city", scope_name="全市（参考件口径）")
        ranks.append(r)

    for i in range(3 if reference else 5):
        dt = (d + timedelta(days=i)).isoformat()
        r = {"id": f"F{i + 1}", "date": dt, "region_id": "demo", "source_id": "S1",
             "issued_at": manifest["sources"][0]["issued_at"], "quality": "valid", "reason": "",
             "basis": "参考件文字预报转录" if reference else "合成预报数据，仅验证展示",
             "aqi_min": [46, 47, 45][i % 3], "aqi_max": [76, 77, 75][i % 3], "category": "优至良",
             "primary_pollutants": ["o3", "pm10"], "weather": "多云" if i == 0 else "晴",
             "temp_min": [17, 16, 17][i % 3], "temp_max": [28, 27, 28][i % 3],
             "humidity_min": [66, 65, 63][i % 3], "humidity_max": [77, 74, 74][i % 3],
             "wind_direction": "东北风", "wind_min": 3 if i < 2 else 2, "wind_max": 4 if i < 2 else 3}
        forecasts.append(r)
    if missing:
        forecasts = []
    for s in manifest["stations"]:
        if not s["focus"]:
            continue
        advice.append({"id": s["id"] + "_advice", "entity_id": s["id"], "source_id": "S1",
                       "issued_at": manifest["sources"][0]["issued_at"], "quality": "valid", "reason": "",
                       "text": "常态化管控。" if reference else "安排现场核查，并由业务人员复核统计资料。",
                       "evidence_ids": []})
    if missing:
        advice = []
    if not missing and not reference:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from airmore.regional.render import preflight
        runtime = preflight()
        for i in range(2):
            with plt.rc_context({"font.family": runtime["font_family"]}):
                fig, ax = plt.subplots(figsize=(6, 3.5), dpi=140)
                ax.plot([0, 6, 12, 18], [1 + i, 2 + i, 4 + i, 2 + i], color="#497D59", marker="o")
                ax.set(xlabel="小时", ylabel="风速 m/s", title=f"合成气象示意图 {i + 1}（仅用于测试）", ylim=(0, 7))
                fig.tight_layout()
                filename = f"weather-{i + 1}.png"
                fig.savefig(root / filename)
                plt.close(fig)
            manifest["weather_images"].append({"id": f"W{i + 1}", "path": filename,
                "region_id": "demo", "source_id": "S1", "caption": "合成气象示意图，仅测试图片接入",
                "issued_at": manifest["sources"][0]["issued_at"],
                "valid_start": (d + timedelta(days=i)).isoformat() + "T00:00:00+08:00",
                "valid_end": (d + timedelta(days=i + 1)).isoformat() + "T00:00:00+08:00"})
    write_json(root / "manifest.json", manifest)
    for name, rows, columns in (("statistics", stats, STAT_COLUMNS), ("ranks", ranks, RANK_COLUMNS),
                                ("contributions", contributions, CONTRIBUTION_COLUMNS)):
        write_csv(root / f"{name}.csv", rows, columns)
    write_json(root / "forecasts.json", forecasts)
    write_json(root / "advice.json", advice)
    return root / "manifest.json"


def make_template(root: Path):
    path = make_example(root, count=1, missing=True)
    m = json.loads(path.read_text("utf-8"))
    m["data_kind"] = "business_import"
    m["sources"][0].update(kind="business_import", label="待填写业务数据来源")
    m["issue"] = "DRAFT"
    for name, columns in (("statistics", STAT_COLUMNS), ("ranks", RANK_COLUMNS),
                          ("contributions", CONTRIBUTION_COLUMNS)):
        write_csv(root / f"{name}.csv", [], columns)
    write_json(path, m)
    return path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--count", type=int, default=4)
    parser.add_argument("--missing", action="store_true")
    parser.add_argument("--reference", type=Path)
    parser.add_argument("--template", action="store_true")
    args = parser.parse_args()
    print(make_template(args.out) if args.template else
          make_example(args.out, args.count, args.missing, reference=args.reference))
