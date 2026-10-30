"""Fill a retained, already reconciled reference report with grounded content.

No data acquisition, model fallback, database writes or service deployment.
Public regional material requires an explicit verified-region evidence package.
"""
import argparse
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

METRICS = ('so2', 'no2', 'co', 'o3', 'pm10', 'pm2_5')
LABELS = ('SO2', 'NO2', 'CO', 'O3', 'PM10', 'PM2.5')


def text(p, value, size=10.5, bold=False):
    p.clear()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = Pt(15)
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = bold
    pr = p._p.get_or_add_pPr()
    snap = pr.find(qn('w:snapToGrid'))
    if snap is None:
        snap = OxmlElement('w:snapToGrid');pr.append(snap)
    snap.set(qn('w:val'), '0')
    r = p.add_run(value)
    r.font.size = Pt(size)
    r.bold = bold
    r.font.name = 'Microsoft YaHei'
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
    return p


def insert_after(doc, anchor, value, bold=False):
    p = doc.add_paragraph()
    text(p, value, bold=bold)
    anchor.addnext(p._p)
    return p._p


def compose(source, evidence, advice, region, target):
    from airmore.service.grounded_advice import validate
    stats = evidence['statistics']; kind = stats['request']['report_type']
    validate({'recommendations': advice['recommendations']}, advice['facts'], kind)
    if advice.get('composition_mode') != 'gpu_original_prose_with_bound_facts' or advice.get('validation') != 'rules_and_independent_model_review_passed' or advice.get('gpu', {}).get('size_vram', 0) <= 0:
        raise ValueError('Verified GPU-generated advice is required')
    if not region.get('confirmed_region') or not region.get('sources') or not region.get('paragraphs'):
        raise ValueError('Confirmed geographic scope and sourced regional material are required')
    for item in region['sources']:
        if not item.get('url', '').startswith('https://') or not item.get('retrieved_at'):
            raise ValueError('Public source URL and retrieval time are required')
    doc = Document(source)
    tables = doc.tables
    hs = list(stats['hangars'].values())
    if len(tables) != 5 or len(tables[1].rows) != len(hs)+1:
        raise ValueError('Reference report structure does not match evidence')
    # Verify inputs before altering any slot. Native units never get converted here.
    for table, slot in ((tables[1], 'last'), (tables[2], 'period')):
        for row, h in zip(table.rows[1:], hs):
            expected = h['days'][h['current']['end']]['survey'] if slot == 'last' else h['current']['survey']
            for j, m in enumerate(METRICS, 1):
                if abs(float(row.cells[j].text) - expected[m]['value']) > .011:
                    raise ValueError('Source table differs from frozen statistics')
    for p in doc.paragraphs:
        if '空气质量排名表' in p.text:
            text(p, p.text.replace('空气质量排名表', '巡测浓度统计表'), bold=True)
        elif p.text.startswith('注：巡测统计'):
            text(p, '注：四项气体单位为ppm，颗粒物为μg/m³；本表为联调设备采样时段统计。区域公开资料在第三章单列。', 9)
    # Replace unsupported rank fields with useful, directly observed coverage.
    for table, slot in ((tables[1], 'last'), (tables[2], 'period')):
        text(table.rows[0].cells[7].paragraphs[0], '有效\n样本', 9, True)
        text(table.rows[0].cells[8].paragraphs[0], '采样\n日期', 9, True)
        for row, h in zip(table.rows[1:], hs):
            c = h['current']
            count = h['days'][c['end']]['samples'] if slot == 'last' else c['samples']
            days = 1 if slot == 'last' else c['observed_days']
            text(row.cells[7].paragraphs[0], str(count), 9)
            text(row.cells[8].paragraphs[0], str(days), 9)
    heading = next(p for p in doc.paragraphs if '各设备六项污染物贡献率占比分析' in p.text)
    text(heading, '（二）各设备六项污染物前后期对比', 12, True)
    heading.paragraph_format.page_break_before = True
    # Two-row panels retain the gray device header but replace empty pie graphics.
    for i, h in enumerate(hs):
        row = tables[3].rows[2*i+1]
        for height in row._tr.xpath('./w:trPr/w:trHeight'):
            height.getparent().remove(height)
        cell = row.cells[0]
        for element in list(cell._tc):
            if element.tag != qn('w:tcPr'):
                cell._tc.remove(element)
        p = cell.add_paragraph()
        a, b = h['current'], h['previous']
        text(p, f"比较期 {b['start']}—{b['end']}（{b['observed_days']}/{b['days']}日有采样）；本期 {a['start']}—{a['end']}（{a['observed_days']}/{a['days']}日有采样）。", 9)
        t = cell.add_table(rows=1, cols=5)
        t.autofit = False
        widths = [1000, 1050, 1550, 1550, 1550]
        for c, label, width in zip(t.rows[0].cells, ('指标','单位','比较期','本期','差值'), widths):
            c.width = width*635
            text(c.paragraphs[0], label, 9, True)
            sh = OxmlElement('w:shd');sh.set(qn('w:fill'), 'E7E6E6');c._tc.get_or_add_tcPr().append(sh)
        for metric, label in zip(METRICS, LABELS):
            x, y = a['survey'][metric], b['survey'][metric]
            if x['unit'] != y['unit'] or x['value'] is None or y['value'] is None:
                raise ValueError('Period comparison requires confirmed identical units and actual values')
            delta = x['value'] - y['value']
            vals = [label, x['unit'], f"{y['value']:.2f}", f"{x['value']:.2f}", f"{delta:+.2f}"]
            for c, value, width in zip(t.add_row().cells, vals, widths):
                c.width = width*635
                text(c.paragraphs[0], value, 9)
        text(cell.paragraphs[-1], '差值为本期减比较期，仅描述原生记录统计；采样条件未完成匹配，不据此认定空气质量改善或恶化。不同单位不相加，不计算六项污染物贡献率。', 9)
    trend = next(p for p in doc.paragraphs if p.text == '采集设备指标变化趋势分析')
    trend.paragraph_format.page_break_before = True
    forecast = next(p for p in doc.paragraphs if p.text == '环境空气质量预测')
    work = next(p for p in doc.paragraphs if p.text == '工作建议')
    # Remove every obsolete forecast paragraph/image, not just visible text.
    cur = forecast._p.getnext()
    while cur is not work._p:
        nxt = cur.getnext();cur.getparent().remove(cur);cur = nxt
    anchor = forecast._p
    anchor = insert_after(doc, anchor, '（一）'+region['confirmed_region']+'区域公开资料', True)
    for paragraph in region['paragraphs']:
        anchor = insert_after(doc, anchor, paragraph)
    for item in region['sources']:
        anchor = insert_after(doc, anchor, '来源：'+item['title']+'；发布／有效时间：'+item['valid_at']+'；'+item['url'])
    # Evidence bindings stay in the audit record; do not repeat entire tables
    # before each original GPU recommendation.
    cur = work._p.getnext()
    while cur is not None and cur.tag != qn('w:sectPr'):
        nxt=cur.getnext();cur.getparent().remove(cur);cur=nxt
    anchor=work._p
    for n,item in enumerate(advice['recommendations']):
        anchor=insert_after(doc,anchor,'（'+'一二三四'[n]+'）'+item['title'],True)
        anchor=insert_after(doc,anchor,item['paragraph'])
    anchor=insert_after(doc,anchor,'统计口径：日内有效分钟等权、期间有采样日期等权；无记录不补零。设备数据来自测试库联调记录；区域公开资料不替代设备实测。正式AQI、综合指数和排名须有符合规范的监测资料。')
    target.parent.mkdir(parents=True,exist_ok=True)
    doc.save(target)
    # python-docx may rewrite unrelated package parts. Preserve all opaque parts.
    with ZipFile(source) as z: original={n:z.read(n) for n in z.namelist()}
    with ZipFile(target) as z: generated={n:z.read(n) for n in z.namelist()}
    original['word/document.xml']=generated['word/document.xml']
    with ZipFile(target,'w',ZIP_DEFLATED) as z:
        for n,v in original.items():z.writestr(n,v)
    return {'file':str(target),'sha256':hashlib.sha256(target.read_bytes()).hexdigest(),
            'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),
            'model':advice['model'],'gpu':advice['gpu'],'sources':region['sources']}


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__)
    for arg in ('source','evidence','advice','region','output'):p.add_argument('--'+arg,type=Path,required=True)
    a=p.parse_args()
    load=lambda x:json.loads(x.read_text())
    print(json.dumps(compose(a.source,load(a.evidence),load(a.advice),load(a.region),a.output),ensure_ascii=False,indent=2))
