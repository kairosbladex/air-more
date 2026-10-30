from datetime import datetime
from decimal import Decimal
import json

import pytest

from airmore.service.hangar_scope import group_selected
from airmore.service.hangar_trend import number, read_snapshot, VERSION
from airmore.service.models import Dataset, METRICS, ReportRequest, ReportError, TZ, survey_value_text
from airmore.service.selection_index import evaluate
from airmore.service.statistics import build_statistics
from tests.test_air_service_worker import Cursor, FakeRepository, fake_render, FixtureWriter
from airmore.service.worker import Worker
from airmore.service.config import Settings


def record(day='2026-09-17', site='s', **values):
    return {'id': 'row', 'project_id': 'p', 'hangar_id': site, 'minute_time': day,
            **{m+'_avg': values.get(m, 1) for m in METRICS}}


def case(kind='DAY', current=None, previous=None, raw=True, sites=('s',)):
    request = ReportRequest(project_id='p', requester_id='reader', hangar_ids=list(sites),
                            start_date={'DAY':'2026-09-17','WEEK':'2026-09-14','MONTH':'2026-09-01'}[kind],
                            end_date='2026-09-17' if kind=='DAY' else '2026-09-20', report_type=kind,
                            as_of=datetime(2026,9,21,tzinfo=TZ))
    samples = [{'id':s,'project_id':'p','hangar_id':s,'time':datetime(2026,9,17,8),
                'load_type':'sensor','load_sn':s,'task_id':'t','lon':None,'lat':None,'alt':None,
                **{m:99 for m in METRICS}} for s in sites] if raw else []
    units=[{'index_name':m,'index_unit':json.dumps([{'productNum':'100301','defaultUnit':'ppm' if m in METRICS[:4] else 'μg/m³'}])} for m in METRICS]
    snapshot={'version':VERSION,'daily':current or [],'periods':{'current':current or [],'previous':previous or []},'series':current or []}
    data=Dataset(samples,{s:s for s in sites},units,provenance={'bindings':[
        {'id':s,'uav_id':'u'+s,'physical_hangar_id':'h','physical_hangar_name':'机库'} for s in sites],
        'registered_payloads':[{'uav_id':'u'+s,'load_brand_code':'100301'} for s in sites]},trend=snapshot)
    return request,data


def statistics(request,data):
    request,data=group_selected(request,data)
    return build_statistics(request,data)


def test_summary_is_authoritative_and_absence_never_falls_back_to_raw():
    expected=dict(so2=.0061,no2=.0487,co=.2926,o3=.0624,pm10=33.3286,pm2_5=29.3333)
    request,data=case(current=[record(**expected),record(**{m:88 for m in METRICS})])
    stats=statistics(request,data);h=stats['hangars']['h']
    assert {m:v['value'] for m,v in h['current']['survey'].items()}==expected
    assert all(v['value'] is None for v in h['previous']['survey'].values())
    assert h['raw_statistics']['current']['survey']['pm10']['value']==99
    assert stats['statistics_source']==VERSION and len(stats['summary_sha256'])==64
    assert h['days']['2026-09-17']['survey']['pm10']['value']==33.3286
    assert survey_value_text(**{k:h['current']['survey']['pm10'][k] for k in ('value','unit','precision')})=='33.3286'
    assert evaluate(stats)['hangars']['h']['previous'] is None


@pytest.mark.parametrize('value,expected', [('1.23445',1.2345),('-1.23445',-1.2345),('0',0), (None,None),('NaN',None)])
def test_java_rounding_signed_values_zero_and_null(value,expected):
    assert number(value)==expected


@pytest.mark.parametrize('kind,key,previous_key', [('WEEK','2026-09-14','2026-09-07'),('MONTH','2026-09-01','2026-08-01')])
def test_week_and_month_use_own_summary_not_daily_average(kind,key,previous_key):
    request,data=case(kind,current=[record(key,pm10=12)],previous=[record(previous_key,pm10=5)])
    data.trend['daily']=[record(pm10=66)]
    h=statistics(request,data)['hangars']['h']
    assert h['current']['survey']['pm10']['value']==12
    assert h['previous']['survey']['pm10']['value']==5
    assert h['days']['2026-09-17']['survey']['pm10']['value']==66
    assert h['current']['summary_periods']==[key]
    assert h['previous']['summary_periods']==[previous_key]


def test_null_and_negative_platform_values_are_not_recomputed_or_converted():
    request,data=case(current=[record(co=None,o3=-.00555,pm10=0)])
    data.unit_configs[4]['index_unit']=json.dumps([{'productNum':'100301','defaultUnit':'mg/m³'}])
    h=statistics(request,data)['hangars']['h']
    assert h['current']['survey']['co']['value'] is None
    assert h['current']['survey']['o3']['value']==-.0056
    assert h['current']['survey']['pm10']['value']==0
    assert h['current']['survey']['pm10']['unit']=='mg/m³'


def test_multi_site_physical_hangar_deduplicates_before_equal_site_mean():
    request,data=case(sites=('a','b'),current=[record(site='a',pm10=4),record(site='a',pm10=99),record(site='b',pm10=8)])
    assert statistics(request,data)['hangars']['h']['current']['survey']['pm10']['value']==6
    data.trend['periods']['current']=[record(site='a',pm10=1.2344),record(site='b',pm10=1.2345)]
    assert statistics(request,data)['hangars']['h']['current']['survey']['pm10']['value']==1.2345
    data.trend['periods']['current']=[record(site='a',pm10=4),record(site='a',pm10=99)]
    assert statistics(request,data)['hangars']['h']['current']['survey']['pm10']['value'] is None


def test_daily_trend_matches_java_summary_not_raw_minutes():
    request,data=case(current=[record()])
    data.trend['series']=[record(o3=-.012345,pm10=Decimal('3.14155'))]
    chart=statistics(request,data)['hangars']['h']['platform_trend']
    series={s['metric']:s['values'] for s in chart['series']}
    assert len(chart['labels'])==15 and chart['labels'][-1]=='2026-09-17'
    assert series['o3'][-1]==-.0123 and series['pm10'][-1]==3.1416
    assert all(v is None for v in series['pm10'][:-1])


def test_disabled_indicator_and_cross_project_records():
    request,data=case(current=[record()])
    data.unit_configs=data.unit_configs[1:]
    assert statistics(request,data)['hangars']['h']['current']['survey']['so2']['value'] is None
    data.trend['periods']['current']=[record()|{'project_id':'another'}]
    with pytest.raises(ReportError,match='超出'):
        statistics(request,data)


def test_units_follow_registered_product_and_ignore_other_device_products():
    request,data=case(current=[record()])
    data.samples[0]['unit_product_num']='100301'
    for unit in data.unit_configs:
        unit['index_unit']=json.dumps([{'productNum':'100301','defaultUnitCode':'101','defaultUnit':''},
                                       {'productNum':'100302','defaultUnit':'mg/m³'}])
        unit['unit_dictionary']=[{'code':'101','value':'ppm'}]
    h=statistics(request,data)['hangars']['h']
    assert h['current']['survey']['co']['unit']=='ppm'
    assert h['current']['survey']['co']['value']==1


def test_scoped_sql_period_tables_tie_breakers_and_boundaries():
    request,_=case('WEEK')
    cur=Cursor([[],[],[],[]]);read_snapshot(cur,request,100)
    assert 'bo_daily_air_report' in cur.sql[0][0]
    assert all('bo_weekly_air_report' in q for q,p in cur.sql[1:])
    assert cur.sql[1][1]==('p','2026-09-14','2026-09-20','s',101)
    assert cur.sql[2][1]==('p','2026-09-07','2026-09-13','s',101)
    for sql,params in cur.sql:
        assert 'project_id=%s AND del_flag=0' in sql
        assert 'COALESCE(update_time,create_time) DESC,id DESC' in sql
    request,_=case()
    cur=Cursor([[],[],[],[]]);read_snapshot(cur,request,100)
    assert 'bo_daily_air_report' in cur.sql[-1][0]
    assert cur.sql[-1][1][1:3]==('2026-09-03','2026-09-17')
    with pytest.raises(ReportError,match='超过'):
        read_snapshot(Cursor([[record(),record()]]),request,1)


def test_worker_can_generate_from_summary_without_raw_rows(tmp_path):
    request,data=case(current=[record()],raw=False)
    worker=Worker(Settings(state_dir=tmp_path),FakeRepository(data),FixtureWriter(),fake_render,
                  forecast_loader=lambda *args:{'gaps':[]})
    job=worker.store.create('backend','summary-only',request)
    worker.run(once=True)
    assert worker.store.owned(job['id'],'backend',request.requester_id)['status'].startswith('succeeded')


def test_render_platform_precision_and_unknown_units_without_crashing(tmp_path,monkeypatch):
    import io
    from PIL import Image
    from docx import Document
    from airmore.service import approved_report
    from airmore.service.approved_advice import write_report
    request,data=case(current=[record(pm10=33.3286,pm2_5=29.3333)])
    data.unit_configs[0]['index_unit']='[]'  # value present, unit unknown
    stats=statistics(request,data)
    image=io.BytesIO();Image.new('RGB',(300,200),'white').save(image,format='PNG')
    for name in ('pie','forecast_card','measured_trend'):
        monkeypatch.setattr(approved_report,name,lambda *a:image.getvalue())
    approved_report.render_document(stats,write_report(None,stats),evaluate(stats),Settings(),tmp_path,{'font_family':'DejaVu Sans'})
    doc=Document(tmp_path/'report.docx')
    assert doc.tables[1].cell(1,5).text=='33.3286'
    assert doc.tables[1].cell(1,6).text=='29.3333'
    assert doc.tables[1].cell(1,1).text=='/'
    assert doc.tables[1].cell(1,7).text=='500'
    assert '未配置' not in '\n'.join(p.text for p in doc.paragraphs)
