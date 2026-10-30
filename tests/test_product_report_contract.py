"""Contracts for the user-approved product report layout."""

from pathlib import Path
from types import SimpleNamespace
from hashlib import sha256

import pytest
from docx import Document
from matplotlib import font_manager

from airmore.regional import product_render
from airmore.regional.product_style import TEMPLATE, TEMPLATE_SHA256, TEMPLATE_VERSION
from airmore.service import product_report
from airmore.service.config import Settings
from airmore.service.examples import synthetic_case
from airmore.service.statistics import build_statistics
from airmore.service.writer import TemplateAnalysisWriter
from airmore.service.models import METRICS


def runtime():
    for family in ("Songti SC", "Noto Serif CJK SC", "Noto Sans CJK SC"):
        try:
            font_path = font_manager.findfont(
                font_manager.FontProperties(family=family), fallback_to_default=False)
            return {"font_family": family,
                    "docx_font_family": "宋体-简" if family == "Songti SC" else family,
                    "font_path": font_path}
        except ValueError:
            continue
    raise RuntimeError("Report render tests require a Chinese font")


def report_case(*, all_missing=False):
    request, data, profiles = synthetic_case(1, missing=True)
    if all_missing:
        for row in data.samples:
            for metric in METRICS:
                row[metric] = None
    stats = build_statistics(request, data, profiles)
    analyses = {
        hid: TemplateAnalysisWriter().write(hangar)
        for hid, hangar in stats["hangars"].items()
    }
    return stats, analyses


def document_text(path):
    doc = Document(path)
    paragraphs = [paragraph.text for paragraph in doc.paragraphs]
    cells = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return "\n".join(paragraphs + cells)


def test_service_and_regional_renderers_share_the_same_template_asset():
    expected_hash = sha256(TEMPLATE.read_bytes()).hexdigest()
    assert product_render.TEMPLATE.resolve() == TEMPLATE.resolve()
    assert product_report.TEMPLATE.resolve() == TEMPLATE.resolve()
    assert TEMPLATE_SHA256 == expected_hash
    assert product_render.TEMPLATE_SHA256 == TEMPLATE_SHA256
    assert product_report.TEMPLATE_SHA256 == TEMPLATE_SHA256
    assert product_report.VERSION == TEMPLATE_VERSION + "+service/v2"
    assert product_render.RENDERER_VERSION == "regional/v1"


def test_environment_specific_settings_do_not_change_layout_identity(tmp_path):
    stats, analyses = report_case()
    results = []
    for name, engine in (("test", "template"), ("production", "ollama")):
        cfg = Settings(state_dir=tmp_path / name, writing_engine=engine)
        output = tmp_path / name / "report"
        results.append(
            product_report.render_document(stats, analyses, {}, cfg, output, runtime())
        )
    assert [result["template_version"] for result in results] == [
        product_report.VERSION,
        product_report.VERSION,
    ]
    assert [result["template_sha256"] for result in results] == [
        TEMPLATE_SHA256,
        TEMPLATE_SHA256,
    ]
    assert all(result["layout"] == "user-approved-product-format" for result in results)


def test_missing_service_metrics_are_rendered_as_slashes(tmp_path):
    stats, analyses = report_case(all_missing=True)
    cfg = Settings(state_dir=tmp_path / "state", writing_engine="template")
    product_report.render_document(stats, analyses, {}, cfg, tmp_path / "report", runtime())
    doc = Document(tmp_path / "report" / "report.docx")
    values = [cell.text for cell in doc.tables[0].rows[1].cells[1:7]]
    assert values == ["/"] * 6
    assert "不可计算" in document_text(tmp_path / "report" / "report.docx")


def test_reference_data_and_source_annotations_do_not_enter_service_report(tmp_path):
    stats, analyses = report_case()
    cfg = Settings(state_dir=tmp_path / "state", writing_engine="template")
    output = tmp_path / "report"
    product_report.render_document(stats, analyses, {}, cfg, output, runtime())
    text = document_text(output / "report.docx")
    for marker in (
        "https://",
        "qwen",
        "GPU",
        "DataV",
        "全市排名125",
        "46-76",
        "20260918",
    ):
        assert marker not in text

    bad_analyses = {
        "demo-1": SimpleNamespace(
            metadata={},
            sections={"overview": "", "advice": "请访问 https://example.invalid/source"},
        )
    }
    with pytest.raises(ValueError, match="source annotation leaked"):
        product_report.render_document(
            stats, bad_analyses, {}, cfg, tmp_path / "bad-report", runtime()
        )


def test_service_table_labels_real_units_and_chinese_period(tmp_path):
    stats, analyses = report_case()
    output = tmp_path / 'report'
    product_report.render_document(stats, analyses, {}, Settings(), output, runtime())
    doc = Document(output / 'report.docx')
    hangar = next(iter(stats['hangars'].values()))
    for metric, cell in zip(METRICS, doc.tables[0].rows[1].cells[1:7]):
        value = hangar['current']['survey'][metric]
        if value['value'] is not None:
            assert value['unit'] in cell.text
    text = document_text(output / 'report.docx')
    assert '报告类型：' + {'DAY': '日报', 'WEEK': '周报', 'MONTH': '月报'}[stats['request']['report_type']] in text
    assert '报告类型：MONTH' not in text


def test_service_entrypoint_renders_selected_site_measurements(monkeypatch, tmp_path):
    from airmore.service import render, regional_context
    stats, analyses = report_case()
    monkeypatch.setenv('AIRMORE_REGIONAL_ENRICHMENT', '1')
    monkeypatch.setattr(render, 'preflight', lambda cfg: runtime())
    monkeypatch.setattr(render, 'convert_pdf', lambda *args: None)
    def reject_city(*args, **kwargs):
        pytest.fail('Site report attempted to fetch city measurements')
    monkeypatch.setattr(regional_context, 'load_context', reject_city)
    render.render(stats, analyses, {}, Settings(), tmp_path)
    doc = Document(tmp_path / 'report.docx')
    rows = doc.tables[2].rows[1:]
    assert len(rows) == len(stats['hangars'])
    for row, site in zip(rows, stats['hangars'].values()):
        assert row.cells[0].text == site['name']
        for metric, cell in zip(METRICS, row.cells[1:7]):
            measurement = site['current']['survey'][metric]
            value, unit = measurement['value'], measurement['unit']
            from airmore.service.models import survey_value_text
            expected = survey_value_text(value, unit) if value is not None and unit else '/'
            assert cell.text == expected
    assert '臭氧为日最大8小时浓度，其余为日均浓度' not in document_text(tmp_path / 'report.docx')
