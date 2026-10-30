from datetime import datetime, timedelta
import json

import pytest

from airmore.service.models import Dataset, ReportRequest
from airmore.service.statistics import build_statistics
from airmore.service.advice import validate_advice


def case(history_days=14):
    request = ReportRequest(hangar_ids=['h'], project_id='p', requester_id='test', report_type='DAY', start_date='2026-09-16', end_date='2026-09-16')
    rows = []
    for i in range(history_days+1):
        stamp = datetime(2026,9,16,9) - timedelta(days=i)
        rows.append(dict(id=str(i), project_id='p', hangar_id='h', load_sn='s', load_type='1003', task_id='t'+str(i),
                         time=stamp, lon=116.0001, lat=36.0001, alt=50, pm10=72 if i==0 else 60, pm2_5=45 if i==0 else 30))
    units = [dict(index_name=m,index_unit=json.dumps([{'defaultUnit':'μg/m³'}])) for m in ('pm10','pm2_5')]
    return request, Dataset(rows, {'h':'测试机库'}, units)


def result(req, data):
    return build_statistics(req, data)['hangars']['h']['relative']


def test_formula_and_change_decomposition():
    r = result(*case())
    assert r['current']['value'] == pytest.approx(135)
    assert r['current']['ratios']['pm2_5'] == pytest.approx(75/135)
    assert r['changes']['net'] == pytest.approx(35)
    assert r['changes']['positive_ratios']['pm2_5'] == pytest.approx(25/35)
    assert r['strata'][0]['reference'] == '30天历史参考'


@pytest.mark.parametrize('count,available', [(0,False),(6,False),(7,True),(13,True)])
def test_history_threshold(count, available):
    r = result(*case(count))
    assert bool(r['current']) == available
    if available:assert r['strata'][0]['reference'] == '短期参考'


@pytest.mark.parametrize('field,value', [('load_sn','other'),('alt',100),('lon',117),('time',datetime(2026,9,16,16))])
def test_unmatched_conditions_do_not_form_reference(field,value):
    req,data=case();data.samples[0][field]=value
    assert result(req,data)['current'] is None


def test_missing_panel_not_reweighted_and_zero_baseline_rejected():
    req,data=case();data.samples[0]['pm2_5']=None
    assert result(req,data)['current'] is None
    req,data=case()
    for row in data.samples[1:]:row['pm10']=0
    assert result(req,data)['current'] is None


def test_no_positive_change_and_no_division_by_zero():
    req,data=case();data.samples[0].update(pm10=30,pm2_5=15)
    r=result(req,data)
    assert r['changes']['positive_ratios'] is None and r['changes']['net'] < 0
    data.samples[0].update(pm10=0,pm2_5=0)
    assert result(req,data)['current']['ratios'] is None


def test_day_balance_does_not_weight_dense_day():
    req,data=case()
    extra=dict(data.samples[1]);extra['pm10']=1000;extra['pm2_5']=1000
    for i in range(50):
        row=extra | {'id':'extra'+str(i), 'time':extra['time']+timedelta(seconds=i+1)}
        data.samples.append(row)
    assert result(req,data)['current']['value'] == pytest.approx(135)


def test_month_to_date_and_cutoff_are_aligned():
    req,_=case()
    req.report_type='MONTH';req.start_date=req.start_date.replace(day=1)
    req.as_of=datetime.fromisoformat('2026-09-16T12:00:00+08:00')
    assert tuple(map(str,req.comparison())) == ('2026-08-01','2026-08-16')
    assert req.comparison_bounds()[1] == datetime(2026,8,16,12)
    assert req.query_bounds()[0] == datetime(2026,8,1)


def test_asof_excludes_future_and_previous_afternoon():
    req,data=case()
    req.as_of=datetime.fromisoformat('2026-09-16T10:00:00+08:00')
    data.samples.append(data.samples[0] | {'id':'later','time':datetime(2026,9,16,11),'pm10':999})
    data.samples.append(data.samples[1] | {'id':'prevlater','time':datetime(2026,9,15,11),'pm10':999})
    h=build_statistics(req,data)['hangars']['h']
    assert h['current']['samples']==1 and h['previous']['samples']==1
    assert h['relative']['current']['value']==pytest.approx(135)


def test_advice_rejects_unknown_evidence_numbers_cause_and_vague_action():
    cards={'coverage':'覆盖不足','level':'浓度可用'}
    item={'evidence_id':'coverage','action':'核对本期缺测时段及原始记录','purpose':'用于核对现有资料能否支持期间对照','verification':'将核对记录整理后提交值班会商复核'}
    body={'recommendations':[item,item | {'evidence_id':'level'}]}
    assert len(validate_advice(body,cards))==2
    for patch in [{'evidence_id':'fake'}, {'action':'复测发现浓度上升50%'}, {'action':'加强监测并提高巡测工作成效'}, {'purpose':'确定污染源为本地工业排放增加'}]:
        with pytest.raises(ValueError):validate_advice({'recommendations':[item | patch,body['recommendations'][1]]},cards)


def test_advice_semantic_rejection_and_repair_are_retained():
    from airmore.service.advice import generate_advice
    from airmore.service.config import Settings
    from airmore.service.models import ReportError
    req,data=case();h=build_statistics(req,data)['hangars']['h']
    a={'evidence_id':'level_pm10','action':'核对本期缺测时段及原始记录','purpose':'用于核对现有资料能否支持期间对照','verification':'将核对记录整理后提交值班会商复核'}
    b=a | {'evidence_id':'relative'}
    class Writer:
        cfg=Settings()
        def _request(self,payload):
            body={'valid':False,'reason':'建议重复'} if 'valid' in payload['format'].get('required',[]) else {'recommendations':[a,b]}
            return {'done':True,'done_reason':'stop','model':self.cfg.ollama_model,'message':{'content':json.dumps(body)}}
        def runtime_probe(self):return {'size_vram':123,'digest':'fixture'}
    writer=Writer()
    with pytest.raises(ReportError,match='未通过事实核验'):generate_advice(writer,h)
    assert len(writer.advice_attempts)==2
    assert all('audit_response' in a and 'validation_error' in a for a in writer.advice_attempts)


def test_full_month_uses_full_previous_calendar_month():
    req,_=case();req.report_type='MONTH';req.start_date=req.start_date.replace(day=1);req.end_date=req.end_date.replace(day=30)
    assert tuple(map(str,req.comparison())) == ('2026-08-01','2026-08-31')


def test_period_render_uses_real_values_and_advice_without_placeholder_charts(tmp_path):
    from airmore.service.period_template import render_document
    from airmore.service.writer import TemplateAnalysisWriter
    from airmore.service.composite import SurveyRelativeComposite
    from airmore.service.config import Settings
    from docx import Document
    req,data=case();stats=build_statistics(req,data)
    analyses={k:TemplateAnalysisWriter().write(v) for k,v in stats['hangars'].items()}
    render_document(stats,analyses,SurveyRelativeComposite().evaluate(stats),Settings(),tmp_path,{'font_family':'Songti SC'})
    doc=Document(tmp_path/'report.docx')
    text='\n'.join(p.text for p in doc.paragraphs)
    assert '135.00' in text and '模型研判待生成' in text
    assert text.index('本班及下一轮观测') < text.index('浓度与观测覆盖')
    assert not doc.inline_shapes
    assert '业务口径待补齐' not in text


def test_advice_accepts_verification_action_synonym():
    item={'evidence_id':'a','action':'核验颗粒物浓度单位及位置高度','purpose':'便于整理同设备同条件历史记录','verification':'整理现有资料标出信息缺口，提交核查清单'}
    assert len(validate_advice({'recommendations':[item,item | {'evidence_id':'b'}]}, {'a':'依据','b':'依据'})) == 2


def test_historical_sensor_report_discloses_missing_link_and_partial_coverage(tmp_path):
    from airmore.service.period_template import render_document
    from airmore.service.writer import TemplateAnalysisWriter
    from airmore.service.config import Settings
    from docx import Document
    req, data = case()
    data.data_kind = 'historical_sensor_replay'
    stats = build_statistics(req, data)
    for h in stats['hangars'].values():
        h['relative']['coverage'] = .005
    analyses = {k: TemplateAnalysisWriter().write(h) for k, h in stats['hangars'].items()}
    render_document(stats, analyses, {}, Settings(), tmp_path, {'font_family': 'Songti SC'})
    text = '\n'.join(p.text for p in Document(tmp_path/'report.docx').paragraphs)
    assert '无法确认历史机库归属' in text
    assert '覆盖率仅0.50%' in text and '不能概括整期' in text
    assert '有效关联采样' not in text


@pytest.mark.parametrize('kind', ['DAY', 'WEEK', 'MONTH'])
def test_full_date_coverage_does_not_prompt_missing_date_advice(kind):
    from airmore.service.advice import generate_advice
    from airmore.service.config import Settings
    req, data = case()
    h = build_statistics(req, data)['hangars']['h']
    h['report_type'] = kind
    class Capture:
        cfg = Settings()
        def _request(self, payload):
            self.payload = payload
            raise RuntimeError('captured')
    writer = Capture()
    with pytest.raises(RuntimeError, match='captured'):
        generate_advice(writer, h)
    prompt = json.loads(writer.payload['messages'][1]['content'])
    assert 'coverage' not in prompt['事实']
    assert '无记录日期' not in json.dumps(prompt, ensure_ascii=False)


def test_native_gas_precision_matches_fact_catalog_and_advice_cards():
    from airmore.service.relative import evidence_cards
    from airmore.service.writer import fact_catalog
    from airmore.service.models import survey_value_text
    req, data = case()
    h = build_statistics(req, data)['hangars']['h']
    h['current']['survey']['so2'].update(value=.008, unit='ppm', days=1)
    facts, _ = fact_catalog(h)
    cards = evidence_cards(h)
    displayed = survey_value_text(.008, 'ppm')
    assert displayed == '0.0080'
    assert any('SO₂' in s and displayed + ' ppm' in s for s in facts.values())
    assert displayed + ' ppm' in next(c['text'] for c in cards if c['id'] == 'level_so2')
