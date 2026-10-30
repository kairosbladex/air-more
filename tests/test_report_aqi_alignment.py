from datetime import date
import io
import json

import pytest
from docx import Document
from PIL import Image

from airmore.service.aqi import calculate, iaqi
from airmore.service.hangar_trend import window, read_snapshot
from airmore.service.indicator_catalog import build_catalog
from airmore.service.report_index import evaluate
from airmore.service.models import METRICS
from tests.test_hangar_trend_alignment import case, record, statistics
from tests.test_air_service_worker import Cursor


@pytest.mark.parametrize('kind,end,first,last,n', [
    ('DAY','2024-03-01','2024-02-16','2024-03-01',15),
    ('WEEK','2026-01-04','2025-11-10','2025-12-29',8),
    ('MONTH','2026-01-31','2025-02-01','2026-01-01',12),
])
def test_calendar_windows(kind,end,first,last,n):
    axis=window(kind,date.fromisoformat(end))
    assert len(axis['keys'])==n and axis['keys'][0]==first and axis['keys'][-1]==last
    if kind=='WEEK':assert axis['labels'][-1]=='2025-12-29~2026-01-04'


@pytest.mark.parametrize('metric,concentration,expected', [
    # Independently transcribed HJ 633—2026 table 3 anchors and interpolations.
    ('so2',150,100),('so2',151,101),('no2',80,100),('no2',81,101),
    ('co','4.05',100),('co','4.15',101),('o3',160,100),('o3',161,101),
    ('pm10',120,100),('pm10',121,101),('pm2_5',60,100),('pm2_5',61,101),
    ('pm2_5','34.5',49),('pm2_5','35.5',52),('o3',801,300),('pm10',601,500),
])
def test_standard_table_values_rounding_and_missing_threshold(metric,concentration,expected):
    assert iaqi(metric,concentration)==expected
    result=calculate({metric:concentration})
    assert result['aqi']==(expected if expected>100 else None)


def test_shared_catalog_masks_standard_aqi_and_platform_index_independently():
    req,data=case(current=[record(so2=150,no2=80,co=4,o3=160,pm10=120,pm2_5=60)],raw=False)
    from airmore.service.models import UNITS
    data.unit_configs=[config(m,UNITS[m]) for m in METRICS]
    stats=statistics(req,data)
    assert stats['hangars']['h']['days']['2026-09-17']['standard']['aqi'] is None
    assert evaluate(stats)['hangars']['h']['current']['value']==100
    data.unit_configs=[r for r in data.unit_configs if r['index_name']!='so2']
    next(r for r in data.unit_configs if r['index_name']=='co')['index_unit']='[]'
    stats=statistics(req,data)
    index=evaluate(stats)['hangars']['h']['current']
    assert set(index['ratios'])=={'no2','o3','pm10','pm2_5'}
    assert index['value'] is None and sum(index['ratios'].values())==pytest.approx(1)
    assert index['excluded']['so2']=='indicator_not_enabled'
    assert index['excluded']['co']=='payload_unit_missing'


def config(metric,unit='ppm',product='100301'):
    return {'index_name':metric,'index_unit':json.dumps([{'productNum':product,'defaultUnit':unit}])}


def test_indicator_presence_and_units_have_distinct_meanings_without_raw_samples():
    req,data=case(current=[record()],raw=False)
    data.unit_configs=[config('SO₂'),config('co',product='wrong'),config('pm2_5','μg/m³')]
    stats=statistics(req,data);h=stats['hangars']['h']
    assert stats['indicator_catalog']['table_metrics']==['so2','co','pm2_5']
    assert h['current']['survey']['so2']['value']==1
    assert h['current']['survey']['co']['value'] is None
    assert h['indicator_validity']['co']['reason']=='payload_unit_missing'
    assert not h['indicator_validity']['pm10']['enabled']
    assert h['current']['summary_available']


@pytest.mark.parametrize('entries,dictionary,reason', [
    ([{'defaultUnit':'ppm'}],[], 'payload_unit_missing'),
    ([{'productNum':'100301','defaultUnitCode':'x'}],[], 'unit_dictionary_conflict'),
    ([{'productNum':'100301','defaultUnitCode':'x','defaultUnit':'ppm'}],
     [{'code':'x','value':'mg/m³'}], 'unit_dictionary_conflict'),
    ([{'productNum':'100301','defaultUnit':'ppm'}, {'productNum':'100301','defaultUnit':'mg/m³'}],[], 'unit_conflict'),
])
def test_units_do_not_fall_back_or_ignore_conflicts(entries,dictionary,reason):
    _,data=case()
    data.unit_configs=[{'index_name':'co','index_unit':json.dumps(entries),'unit_dictionary':dictionary}]
    c=build_catalog(data)
    assert c['hangars']['h']['co']=={'enabled':True,'valid':False,'unit':None,'reason':reason}


def test_multiple_registered_products_require_consistent_units():
    req,data=case(current=[record()])
    data.provenance['registered_payloads'].append({'uav_id':'us','load_brand_code':'100302'})
    h=statistics(req,data)['hangars']['h']
    assert all(v['value'] is None for v in h['current']['survey'].values())


def test_unit_gate_cannot_be_overridden_by_measurement_profiles():
    from airmore.service.statistics import measurement
    row={'load_type':'sensor','co':5}
    profile={'basis':'confirmed','units':{'co':{'unit':'mg/m³','confirmed':True}}}
    validity={'valid':False,'reason':'payload_unit_missing'}
    assert measurement(row,'co',profile,{},validity)==(None,None,None,'payload_unit_missing')


def test_trend_extra_indicators_and_methane_do_not_expand_aqi():
    req,data=case('WEEK',current=[record('2026-09-14')])
    data.unit_configs += [config('PM1.0','μg/m³'),config('OX'),config('CH₄')]
    data.trend['series']=[record('2026-09-14')|{'pm1_0_avg':4,'ox_avg':0}]
    data.trend['methane']=[record('2026-09-07')|{'ch4_avg':3},record('2026-09-14')|{'ch4_avg':2}]
    stats=statistics(req,data);chart=stats['hangars']['h']['platform_trend']
    values={s['metric']:s['values'] for s in chart['series']}
    assert values['ch4'][-2:]==[3,2] and values['pm1_0'][-1]==4 and values['ox'][-1]==0
    assert stats['indicator_catalog']['table_metrics']==list(METRICS)
    assert 'ch4' not in calculate({'ch4':999})['iaqi']


def test_methane_sql_is_independent_and_scoped():
    req,_=case('MONTH');cur=Cursor([[],[],[],[],[]])
    read_snapshot(cur,req,100,['co','ch4'])
    sql,args=cur.sql[-1]
    assert 'bo_task_methane' in sql and 'bo_task_air' not in sql
    assert 'md.concentration>=0' in sql and 'AND t.project_id=%s AND s.project_id=%s' in sql
    assert args[0]=='p' and args[-2:] == ('s',101)
    assert args[1].date()==date(2025,10,1) and args[2].date()==date(2026,9,20)


@pytest.mark.parametrize('kind,table_count', [('DAY',1),('WEEK',2),('MONTH',2)])
def test_document_tables_dynamic_columns_slashes_and_removed_notes(kind,table_count,tmp_path,monkeypatch):
    from airmore.service import approved_report as report
    from airmore.service.approved_advice import write_report
    from airmore.service.config import Settings
    req,data=case(kind,current=[record('2026-09-17' if kind=='DAY' else '2026-09-14')])
    data.unit_configs=[config('so2'),config('co',product='other')]
    data.trend['daily']=[record(req.end_date.isoformat(),so2=7)]
    stats=statistics(req,data)
    img=io.BytesIO();Image.new('RGB',(20,20),'white').save(img,format='PNG')
    for name in ('pie','forecast_card','measured_trend'):monkeypatch.setattr(report,name,lambda *args:img.getvalue())
    report.render_document(stats,write_report(None,stats),evaluate(stats),Settings(),tmp_path,{'font_family':'DejaVu Sans'})
    doc=Document(tmp_path/'report.docx');tables=doc.tables[1:1+table_count]
    assert len(doc.tables)==table_count+4  # Includes the advice comparison table.
    assert [c.text for c in tables[0].rows[0].cells]==['机库名称','SO2','CO','综合指数','名次']
    assert tables[0].cell(1,1).text=='7.0000' and tables[0].cell(1,2).text=='/'
    assert all(len(t.rows[0].cells)==5 for t in tables)
    text='\n'.join(p.text for p in doc.paragraphs)
    assert '综合相对指数' not in text and '污染物分指数占比分析' in text
    full_text='\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)+text
    assert all(x not in full_text for x in ('空白表示', '日AQI', '平台日平均值', '位置断线', '有效天数'))
