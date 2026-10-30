import json
from datetime import datetime

import pytest
from docx import Document
from tests.test_air_service_relative import case
from airmore.service.statistics import build_statistics
from airmore.service.report_facts import fact_pack, series
from airmore.service.report_advice import validate, generate_advice
from airmore.service.models import ReportError
from airmore.service.config import Settings


def test_hourly_survey_uses_equal_minutes_and_preserves_gaps():
    req,data=case()
    rows=[r for r in data.samples if r['time'].date()==req.end_date]
    row=rows[0]
    data.samples=[row | {'id':'a','time':datetime(2026,9,16,9,0,0),'pm10':10},
                  row | {'id':'b','time':datetime(2026,9,16,9,0,30),'pm10':30},
                  row | {'id':'c','time':datetime(2026,9,16,9,1,0),'pm10':100},
                  row | {'id':'d','time':datetime(2026,9,16,11,0,0),'pm10':200}]
    h=build_statistics(req,data)['hangars']['h']
    assert h['days']['2026-09-16']['hourly']['09:00']['pm10']['value']==60
    points,basis=series(h)
    assert basis=='分钟巡测均值'
    assert dict(points)['10:00']=={}
    assert h['current']['survey']['pm10']['value']==pytest.approx((20+100+200)/3)
    assert '11:00' in fact_pack(h)['peak']


def advice_items():
    return [{'evidence_id':'peak','title':'高值时段对照观测','paragraph':'下一轮观测优先覆盖该时段，在相近位置和记录高度开展复测，并与可取得的同期地面站记录对照。对照时区分采样时间和空间范围，保留颗粒物变化过程供值班研判。'},
            {'evidence_id':'comparison','title':'保持期间对比条件','paragraph':'后续观测保留同设备同位置的固定时段和高度，逐步增加可匹配的历史记录。会商分析应区分局部可比观测与全部采样结果，避免将局部指数变化外推至整个区域。'}]


def test_report_prose_requires_valid_evidence_and_rejects_invented_numbers():
    facts={'peak':'峰值事实','comparison':'可比事实'}
    items=advice_items()
    assert validate({'recommendations':items},facts,'DAY')==items
    for patch in ({'evidence_id':'unknown'},{'paragraph':items[0]['paragraph']+'浓度增加50%。'},
                  {'paragraph':items[0]['paragraph']+'我认为应提交核查清单。'},
                  {'paragraph':items[0]['paragraph']+'下月调整观测安排。'}):
        with pytest.raises(ValueError):validate({'recommendations':[items[0]|patch,items[1]]},facts,'DAY')


def test_report_prose_retains_failed_semantic_review():
    req,data=case();h=build_statistics(req,data)['hangars']['h']
    class Writer:
        cfg=Settings()
        def _request(self,payload):
            content={'valid':False,'reason':'语义不成立'} if 'valid' in payload['format']['required'] else {'selections':['peak_1','comparison_1']}
            return {'done':True,'done_reason':'stop','model':self.cfg.ollama_model,'message':{'content':json.dumps(content)}}
    writer=Writer()
    with pytest.raises(ReportError):generate_advice(writer,h)
    assert len(writer.advice_attempts)==2
    assert all('audit_response' in r for r in writer.advice_attempts)


def test_report_prose_success_binds_facts_and_keeps_gpu_audit():
    req,data=case();h=build_statistics(req,data)['hangars']['h']
    class Writer:
        cfg=Settings()
        def _request(self,payload):
            content={'valid':True,'reason':'事实相符'} if 'valid' in payload['format']['required'] else {'selections':['peak_1','comparison_1']}
            return {'done':True,'done_reason':'stop','model':self.cfg.ollama_model,'message':{'content':json.dumps(content)}}
        def runtime_probe(self):return {'size_vram':123,'digest':'fixture'}
    prose,audit=generate_advice(Writer(),h)
    assert fact_pack(h)['peak'] in prose and fact_pack(h)['comparison'] in prose
    assert audit['gpu']['digest']=='fixture'
    assert len(audit['attempts'])==1 and len(audit['paragraphs'])==2
    assert audit['paragraphs'][0]['paragraph'].endswith(audit['choices']['peak_1']['paragraph'])
    assert audit['composition_mode']=='gpu_selection_of_evidence_specific_business_prose'


def test_polished_report_keeps_facts_and_removes_empty_charts(tmp_path,monkeypatch):
    from airmore.service.report_template import render_document
    from airmore.service.writer import TemplateAnalysisWriter
    req,data=case();stats=build_statistics(req,data)
    monkeypatch.setattr('airmore.service.report_template.trend_image',lambda *a:None)
    analyses={k:TemplateAnalysisWriter().write(h) for k,h in stats['hangars'].items()}
    render_document(stats,analyses,{},Settings(),tmp_path,{'font_family':'Songti SC'})
    doc=Document(tmp_path/'report.docx');body='\n'.join(p.text for p in doc.paragraphs)
    assert len(doc.tables)==2
    assert len(doc.tables[0].rows)==7
    assert doc.tables[0].cell(0,1).text=='单位'
    assert '核查清单' not in body and '未来五天AQI数值预报' not in body
    assert body.index('一、空气质量情况')<body.index('二、采集指标变化趋势')<body.index('四、工作建议')
