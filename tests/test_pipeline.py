"""端到端流水线测试：模拟数据 → 9 个产物 → DOCX 读回抽查。

需要 soffice 进行真实 PDF 转换；缺失时跳过。
"""

import importlib.util
import json
import shutil
from pathlib import Path

import pytest
from docx import Document

from airmore.pipeline import run

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SOFFICE = "/opt/homebrew/bin/soffice"

pytestmark = pytest.mark.skipif(
    not (Path(SOFFICE).exists() or shutil.which("soffice")),
    reason="soffice 不可用，跳过端到端 PDF 测试",
)


def _generate_sample(data_dir: Path) -> None:
    spec = importlib.util.spec_from_file_location(
        "gen_sample", PROJECT_ROOT / "scripts" / "generate_sample_data.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    rng = mod.np.random.default_rng(mod.SEED)
    data_dir.mkdir(parents=True, exist_ok=True)
    df = mod.pd.concat(
        [mod.build_sortie(sid, hid, s, e, step, rng)
         for sid, hid, s, e, step in mod.SORTIES],
        ignore_index=True)
    df = mod.inject_defects(df).sort_values("timestamp").reset_index(drop=True)
    for hid in mod.HANGARS:
        df[df["hangar_id"] == hid].to_csv(
            data_dir / f"{hid}_{mod.DAY}.csv", index=False)


@pytest.fixture(scope="module")
def pipeline_result(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("e2e")
    data_dir = tmp / "data"
    _generate_sample(data_dir)
    out_dir = tmp / "out"
    result = run(data_dir, out_dir, ["A", "B", "G"])
    return result, out_dir


def _docx_all_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)


def test_nine_artifacts_exist_and_nonempty(pipeline_result):
    result, out_dir = pipeline_result
    assert set(result["artifacts"]) == {"A", "B", "G"}
    for art in result["artifacts"].values():
        for key in ("json", "docx", "pdf"):
            p = Path(art[key])
            assert p.exists() and p.stat().st_size > 1000, p
    # 贡献率饼图等新图表
    assert (out_dir / "charts" / "contribution_pies.png").exists()


def test_pdf_pages_reasonable(pipeline_result):
    result, _ = pipeline_result
    pages = {k: v["pdf_pages"] for k, v in result["artifacts"].items()}
    assert 3 <= pages["A"] <= 4   # 目标约 3 页，允许版式浮动
    assert 2 <= pages["B"] <= 3   # 目标约 2 页
    assert 2 <= pages["G"] <= 4   # 目标约 2–3 页


def test_pm25_event_detected(pipeline_result):
    result, _ = pipeline_result
    assert result["events"] >= 1
    b = json.loads(Path(result["artifacts"]["B"]["json"]).read_text("utf-8"))
    event_items = [e for e in b["evidence"] if e["kind"] == "event"]
    assert any(e["extra"]["sortie_id"] == "S-A-2" and e["description"].startswith("PM2.5")
               for e in event_items)


def test_a_report_govt_header_and_structure(pipeline_result):
    result, _ = pipeline_result
    doc = Document(result["artifacts"]["A"]["docx"])
    text = _docx_all_text(doc)
    # 三行报头
    assert "无人机低空环境空气质量研判日报" in text
    assert "第﹝2026﹞第001期（模拟期号）" in text
    assert "2026年3月15日" in text
    # 四个固定节 + 附注
    for section in ["一、巡测区域空气质量情况", "二、环境空气质量预测",
                    "三、空气污染气象条件", "四、工作建议"]:
        assert section in text
    # 预测明确不可用，不编造 AQI
    assert "预测不可用" in text
    assert "无数据（无外部预报接入）" in text
    # 口径与演示声明
    assert "非全天 24h 均值" in text
    assert "非正式达标判定" in text
    assert "排名后三字体为黄色，排名最后一位字体为红色" in text
    # 当月表不编造累计
    assert "暂不累计" in text
    # 工作建议含排名、事件证据引用与管控措辞
    assert "常态化管控" in text
    assert "业务筛查阈值不等于法定超标判定" in text


def _find_table(doc: Document, first_header: str):
    for t in doc.tables:
        if t.rows[0].cells[0].text.strip() == first_header:
            return t
    raise AssertionError(f"未找到表头为 {first_header} 的表")


def test_a_stats_table_values_match_evidence(pipeline_result):
    """巡测日统计表逐格与证据 JSON 一致（数值可复算）。"""
    result, _ = pipeline_result
    pkg = json.loads(Path(result["artifacts"]["A"]["json"]).read_text("utf-8"))
    ev = {e["description"].split(" ")[0]: e
          for e in pkg["evidence"] if e["kind"] == "hangar_index"}
    assert len(ev) == 4
    doc = Document(result["artifacts"]["A"]["docx"])
    table = _find_table(doc, "2026.3.15")

    fields = ["so2", "no2", "co", "o3", "pm10", "pm2_5"]
    for row in table.rows[1:]:
        cells = [c.text.strip() for c in row.cells]
        hid = cells[0]
        e = ev[hid]
        means = e["extra"]["means"]
        for j, f in enumerate(fields, start=1):
            v = means.get(f)
            expected = "/" if v is None else (f"{v:.1f}" if f == "co" else f"{v:.0f}")
            assert cells[j] == expected, (hid, f, cells[j], expected)
        assert cells[7] == f"{e['extra']['index']:.2f}"
        assert cells[8] == str(e["extra"]["rank"])

    # 排名颜色：后三黄色、末位红色
    from docx.shared import RGBColor
    n = len(table.rows) - 1
    for i, row in enumerate(table.rows[1:], start=0):
        rank = int(row.cells[8].text.strip())
        colors = {r.font.color.rgb for c in row.cells for p in c.paragraphs
                  for r in p.runs if r.font.color and r.font.color.rgb}
        if rank == n:
            assert RGBColor(0xFF, 0x00, 0x00) in colors, (i, rank)
        elif rank > n - 3:
            assert RGBColor(0xFF, 0xC0, 0x00) in colors, (i, rank)
        else:
            assert not colors, (i, rank, colors)

    # 当月累计表全为 "/"
    monthly = _find_table(doc, "2026年3月")
    for row in monthly.rows[1:]:
        assert all(c.text.strip() == "/" for c in row.cells[1:])


def test_docx_readback_disclaimer_and_numbers(pipeline_result):
    from airmore.validate.numbers import _EVIDENCE_REF, _PATTERN

    for key, stem in [("A", "A_巡测日报"), ("B", "B_异常快报"), ("G", "G_专项建议")]:
        result, _ = pipeline_result
        doc = Document(result["artifacts"][key]["docx"])
        text = _docx_all_text(doc)
        footer = "\n".join(p.text for s in doc.sections for p in s.footer.paragraphs)
        assert "模拟数据演示" in footer

        pkg = json.loads(Path(result["artifacts"][key]["json"]).read_text("utf-8"))
        allowed = [(float(n["value"]), n["unit"])
                   for e in pkg["evidence"] for n in e["numbers"]]
        ids = {e["id"] for e in pkg["evidence"]}
        # 全文每一个"数值+单位"都能在证据包找到
        for raw, unit in _PATTERN.findall(text):
            assert any(abs(float(raw) - v) <= 0.051 and u == unit
                       for v, u in allowed), (stem, raw, unit)
        # 全文每一个证据引用都存在
        assert set(_EVIDENCE_REF.findall(text)) <= ids

    # A 类：工作建议中的事件与综合指数证据编号在正文出现
    pkg_a = json.loads(Path(result["artifacts"]["A"]["json"]).read_text("utf-8"))
    text_a = _docx_all_text(Document(result["artifacts"]["A"]["docx"]))
    event_ids = [e["id"] for e in pkg_a["evidence"] if e["kind"] == "event"]
    index_ids = [e["id"] for e in pkg_a["evidence"] if e["kind"] == "hangar_index"]
    for eid in event_ids + index_ids:
        assert f"[{eid}]" in text_a, eid
    # B/G 报头统一风格
    for key in ("B", "G"):
        btext = _docx_all_text(Document(result["artifacts"][key]["docx"]))
        assert "（模拟期号）" in btext
        assert "2026年3月15日" in btext


def test_json_contains_disclaimer_and_rules(pipeline_result):
    result, _ = pipeline_result
    for art in result["artifacts"].values():
        pkg = json.loads(Path(art["json"]).read_text("utf-8"))
        assert pkg["disclaimer"] == "模拟数据演示，不用于业务决策"
        assert "仅供演示" in pkg["rule_info"]["disclaimer"]
        assert pkg["narrative"]["_validation"].startswith("通过")
