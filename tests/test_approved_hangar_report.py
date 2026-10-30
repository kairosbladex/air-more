from copy import deepcopy
from dataclasses import replace
from datetime import date, datetime
import io
import json
from pathlib import Path

import pytest
from PIL import Image
from docx import Document

from airmore.service.config import Settings
from airmore.service.models import METRICS, Dataset, ReportError, ReportRequest, TZ
from airmore.service.hangar_scope import group_selected
from airmore.service.selection_index import evaluate, competition_ranks
from airmore.service.statistics import configured_units, measurement
from airmore.service.payloads import select_atmospheric_samples
from airmore.service.site_forecast import digest, validate_snapshot, load


def stats_fixture():
    survey={m:{'value':10.,'unit':'ppm' if m in METRICS[:4] else 'μg/m³'} for m in METRICS}
    hs={key:{'six_relative':{'panel':list(METRICS),'strata':[{'baseline_days':9,'baseline':{m:10. for m in METRICS}}],
                          'units':{m:[v['unit']] for m,v in survey.items()}},
             'current':{'survey':deepcopy(survey)},'previous':{'survey':deepcopy(survey)},
             'days':{'2026-09-18':{'survey':deepcopy(survey)}}} for key in ('a','b','c')}
    return {'request':{'hangar_ids':['a','b','c'],'end_date':'2026-09-18'},'hangars':hs}


def test_selected_index_scope_ties_and_fail_closed():
    stats=stats_fixture();stats['hangars']['c']['current']['survey']['co']['value']=20
    result=evaluate(stats)
    assert {k:r['current']['rank'] for k,r in result['hangars'].items()}=={'a':1,'b':1,'c':3}
    assert all(sum(r['current']['ratios'].values())==pytest.approx(1) for r in result['hangars'].values())
    del stats['hangars']['b'];stats['request']['hangar_ids'].remove('b')
    assert evaluate(stats)['selected_hangar_count']==2
    stats['hangars']['a']['current']['survey']['co']['unit']='mg/m³'
    assert evaluate(stats)['gaps']==[{'code':'SIX_UNIT_MISMATCH'}]
    stats['hangars']['a']['current']['survey']['co']['unit']='ppm'
    stats['hangars']['a']['current']['survey']['co']['value']=None
    partial=evaluate(stats)
    assert partial['available'] and partial['hangars']['a']['current'] is None
    assert competition_ranks({'a':1.001,'b':1.004,'c':None})=={'a':1,'b':1}


def test_physical_hangar_groups_preserve_original_site_scope():
    req=ReportRequest(project_id='p',requester_id='r',hangar_ids=['s1','s2'],start_date='2026-09-18',end_date='2026-09-18')
    data=Dataset([{'hangar_id':'s1','id':'1'},{'hangar_id':'s2','id':'2'}],{'s1':'站点一','s2':'站点二'},
                 provenance={'bindings':[{'id':s,'physical_hangar_id':'h','physical_hangar_name':'真实机库'} for s in req.hangar_ids]})
    scoped,grouped=group_selected(req,data)
    assert scoped.hangar_ids==['h'] and req.hangar_ids==['s1','s2']
    assert grouped.hangars=={'h':'真实机库'}
    assert [r['source_site_id'] for r in grouped.samples]==['s1','s2']
    assert data.samples[0]['hangar_id']=='s1'
    with pytest.raises(ReportError):group_selected(req,replace(data,provenance={}))


def test_dictionary_unit_resolution_follows_registered_payload_and_rejects_conflicts():
    raw={'hangar_id':'s','load_sn':'serial','load_type':'1003','co':.2926}
    payload={'id':'l','uav_id':'u','load_imei':'serial','load_type':'1003','load_brand_code':'100301'}
    rows,_,_=select_atmospheric_samples([raw],[{'id':'s','uav_id':'u'}],[payload],('1003',))
    assert rows[0]['unit_product_num']=='100301' and rows[0]['load_type']=='1003'
    record={'index_name':'co','index_unit':json.dumps([{'productNum':'100301','defaultUnitCode':'101','defaultUnit':''}]),
            'unit_dictionary':[{'code':'101','value':'ppm'}]}
    assert measurement(rows[0],'co',{},configured_units([record]))==(.2926,'ppm',None,'measurement_state_unconfirmed')
    assert measurement(raw,'co',{},configured_units([record]))[3]=='unit_unknown'
    record['unit_dictionary'].append({'code':'101','value':'mg/m³'})
    assert measurement(rows[0],'co',{},configured_units([record]))[3]=='unit_conflict'
    record['unit_dictionary']=[{'code':'101','value':'unknown'}]
    assert measurement(rows[0],'co',{},configured_units([record]))[3]=='unit_conflict'


def archive_fixture():
    from datetime import timedelta
    issue=date(2026,9,18)
    payload={'city_code':'410100','city_name':'郑州市','issue_date':issue.isoformat(),
             'forecasts':[{'areaCode':'410100','dataDate':(issue+timedelta(days=i)).strftime('%Y%m%d'),
                           'aqi_min':50,'aqi_max':80,'primaryPol':'O3'} for i in range(5)],
             'weather':[{'date':(issue+timedelta(days=i)).isoformat(),'high':'26','low':'19','wind':'北风','force':'小于3级','description':'多云'} for i in range(5)],
             'sources':[{'retrieved_at':'2026-09-18T12:00:00+08:00','sha256':'fixture'}]}
    return {'payload':payload,'sha256':digest(payload)}


def test_forecast_archive_scope_date_and_checksum(tmp_path):
    snapshot=archive_fixture();issue=date(2026,9,18)
    assert validate_snapshot(snapshot,'410100',issue)['city_name']=='郑州市'
    for code,day in [('wrong',issue),('410100',date(2026,9,19))]:
        with pytest.raises(ValueError):validate_snapshot(snapshot,code,day)
    snapshot['payload']['weather'][0]['high']='99'
    with pytest.raises(ValueError):validate_snapshot(snapshot,'410100',issue)
    stats={'request':{'end_date':'2026-09-18','hangar_ids':['s'],'project_id':'p',
                      'as_of':'2026-09-18T12:00:00+08:00'}}
    sites=[{'id':'s','name':'设备','position':'郑州市','lng_lat_alt':None}]
    cfg=Settings(state_dir=tmp_path,forecast_enabled=True)
    def no_network(*args):raise AssertionError('historical forecasts must use archives')
    missing=load(stats,cfg,tmp_path/'missing',sites=sites,project={},read=no_network,now=datetime(2026,9,20,tzinfo=TZ))
    assert missing['gaps'] and not missing['forecasts']
    path=tmp_path/'forecast-archive/410100/2026-09-18.json';path.parent.mkdir(parents=True);path.write_text(json.dumps(archive_fixture()))
    loaded=load(stats,cfg,tmp_path/'loaded',sites=sites,project={},read=no_network,now=datetime(2026,9,20,tzinfo=TZ))
    assert len(loaded['forecasts'])==5 and not loaded['gaps']


def test_approved_template_values_names_and_missing_index(tmp_path,monkeypatch):
    from airmore.service.examples import synthetic_case
    from airmore.service.statistics import build_statistics
    from airmore.service.approved_advice import write_report
    from airmore.service import approved_report as render
    req,data,profiles=synthetic_case(1,missing=True)
    stats=build_statistics(req,data,profiles)
    stats['hangars']['demo-1']['name']='已核对机库名称'
    stats['forecast']=archive_fixture()['payload']
    stats['forecast']['weather'][0]['high']=None
    stats['forecast']['weather'][1]['high']='26℃'
    for row in stats['forecast']['forecasts']:
        row['dataDate']=datetime.strptime(row['dataDate'],'%Y%m%d').date().isoformat()
    buffer=io.BytesIO();Image.new('RGB',(300,200),'white').save(buffer,format='PNG')
    for name in ('pie','forecast_card','measured_trend'):monkeypatch.setattr(render,name,lambda *a:buffer.getvalue())
    result=render.render_document(stats,write_report(None,stats),evaluate(stats),Settings(),tmp_path,{'font_family':'DejaVu Sans'})
    doc=Document(tmp_path/'report.docx')
    assert doc.tables[1].cell(0,0).text=='机库名称'
    assert doc.tables[1].cell(1,0).text=='已核对机库名称'
    assert doc.tables[1].cell(1,8).text=='/'
    assert doc.tables[1].cell(1,7).text=='/'
    assert result['scope']=='physical_hangar'
    text=''.join(p.text for p in doc.paragraphs)
    assert '气温最低19℃' in text and '气温19—26℃，' in text
    assert '℃℃' not in text and 'None' not in text
    assert '合成数据测试' in text and '资料说明' not in text
    assert len([p for p in doc.paragraphs if '值班人员' in p.text])==1
