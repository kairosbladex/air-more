"""Regional period/provenance contracts and real artifact regression checks."""

import copy
import csv
import hashlib
import json
import shutil
from pathlib import Path

import pytest
from docx import Document

from airmore.regional.bundle import load_bundle
from airmore.regional.report import build_view, validate_view
from scripts.regional_examples import make_example, write_csv, write_json


@pytest.fixture(scope="module")
def template(tmp_path_factory):
    return make_example(tmp_path_factory.mktemp("regional-template"), missing=True)


@pytest.fixture
def bundle_path(tmp_path, template):
    target = tmp_path / "input"
    shutil.copytree(template.parent, target)
    return target / "manifest.json"


def manifest_edit(path, edit):
    m = json.loads(path.read_text())
    edit(m)
    write_json(path, m)


def csv_edit(path, name, edit):
    f = path.parent / f"{name}.csv"
    with f.open() as stream:
        reader = csv.DictReader(stream)
        columns = reader.fieldnames
        rows = list(reader)
    edit(rows)
    write_csv(f, rows, columns)


def test_month_boundary_and_hash_snapshot(bundle_path):
    b = load_bundle(bundle_path)
    assert b.period("day") == ("2026-08-31", "2026-08-31")
    assert b.period("month") == ("2026-08-01", "2026-08-31")
    assert b.inputs[0]["sha256"] == hashlib.sha256(bundle_path.read_bytes()).hexdigest()
    assert len(b.inputs) == 6


@pytest.mark.parametrize("dt,expected", [("2027-01-01", ("2026-12-01", "2026-12-31")),
                                         ("2028-03-01", ("2028-02-01", "2028-02-29")),
                                         ("2026-09-09", ("2026-09-01", "2026-09-08"))])
def test_calendar_windows(tmp_path, dt, expected):
    b = load_bundle(make_example(tmp_path, missing=True, report_date=dt))
    assert b.period("month") == expected


@pytest.mark.parametrize("mutation", [
    lambda r: r[0].update(start="2026-08-01"),
    lambda r: r[0].update(end="2026-09-01"),
    lambda r: r[0].update(entity_id="OTHER"),
    lambda r: r[0].update(unit="ppm"),
    lambda r: r.append(dict(r[0], id="COPY")),
    lambda r: r[0].update(quality="valid", value="NaN"),
    lambda r: r[0].update(quality="valid", value="-1"),
    lambda r: r[0].update(value="100"),
    lambda r: r[0].update(reason=""),
    lambda r: r[0].update(source_id="UNKNOWN"),
])
def test_invalid_statistics_stop(bundle_path, mutation):
    csv_edit(bundle_path, "statistics", mutation)
    with pytest.raises(ValueError):
        load_bundle(bundle_path)


@pytest.mark.parametrize("mutation", [
    lambda r: r[0].update(quality="valid", rank="5", total="4"),
    lambda r: r[0].update(total="0"),
    lambda r: r[0].update(quality="valid", rank="1.2"),
    lambda r: r[0].update(scope_name="另一个范围"),
])
def test_invalid_rank_scope(bundle_path, mutation):
    csv_edit(bundle_path, "ranks", mutation)
    with pytest.raises(ValueError):
        load_bundle(bundle_path)


def test_preserve_imported_index_missing_not_ranked(bundle_path):
    csv_edit(bundle_path, "statistics", lambda rs: next(r for r in rs if r["metric"] == "composite").update(
        quality="valid", value="2.16", reason=""))
    b = load_bundle(bundle_path)
    v = build_view(b)
    table = next(x for x in v["blocks"] if x["type"] == "table")
    assert table["rows"][0][-2:] == ["2.16", "/"]
    assert table["styles"][0] == {"fill": "FFFF00", "color": None}


def test_missing_sections_and_all_missing_contributions(bundle_path):
    v = build_view(load_bundle(bundle_path))
    assert all(p["ratios"] is None for c in v["charts"] for p in c["panels"])
    text = "\n".join(b.get("text", "") for b in v["blocks"])
    for section in ("二、环境空气质量预测", "三、空气污染气象条件", "四、工作建议"):
        assert section in text
    assert "常态化管控" not in text
    assert v["missing"]


def test_no_naked_number_escape(bundle_path):
    b = load_bundle(bundle_path)
    v = build_view(b)
    table = next(x for x in v["blocks"] if x["type"] == "table")
    table["rows"][0][-2] = "999.99"
    with pytest.raises(ValueError, match="bound source"):
        validate_view(b, v)


def test_contribution_modes_and_zero(bundle_path):
    def zeros(rows):
        for r in rows:
            r.update(value="0", quality="valid", reason="", mode="components")
    csv_edit(bundle_path, "contributions", zeros)
    assert all(p["ratios"] is None for c in build_view(load_bundle(bundle_path))["charts"] for p in c["panels"])
    csv_edit(bundle_path, "contributions", lambda rs: rs[0].update(mode="ratios"))
    with pytest.raises(ValueError, match="mixed contribution"):
        load_bundle(bundle_path)


def test_ratio_sum_invalid(bundle_path):
    csv_edit(bundle_path, "contributions", lambda rs: [r.update(quality="valid", value="0.4", reason="") for r in rs])
    with pytest.raises(ValueError, match="sum to one"):
        load_bundle(bundle_path)


@pytest.mark.parametrize("edit", [
    lambda m: m.update(published_at="2026-09-01T09:00:00"),
    lambda m: m["sources"][0].update(issued_at="2026-09-01T10:00:00+08:00"),
    lambda m: m["sources"][0].update(kind="business_import"),
    lambda m: m["stations"].append(copy.deepcopy(m["stations"][0])),
    lambda m: m["files"].update(statistics="../outside.csv"),
])
def test_manifest_rejection(bundle_path, edit):
    manifest_edit(bundle_path, edit)
    with pytest.raises(ValueError):
        load_bundle(bundle_path)


@pytest.fixture(scope="module")
def complete(tmp_path_factory):
    return make_example(tmp_path_factory.mktemp("regional-full"))


@pytest.fixture
def complete_path(tmp_path, complete):
    shutil.copytree(complete.parent, tmp_path / "full")
    return tmp_path / "full" / "manifest.json"


@pytest.mark.parametrize("field,value", [("date", "2026-08-31"), ("region_id", "OTHER"),
    ("aqi_min", 800), ("aqi_max", 1), ("aqi_min", 1.5), ("humidity_max", 101),
    ("temp_min", None), ("issued_at", "2026-09-01T10:00:00+08:00")])
def test_forecast_validation(complete_path, field, value):
    p = complete_path.parent / "forecasts.json"
    data = json.loads(p.read_text())
    data[0][field] = value
    write_json(p, data)
    with pytest.raises(ValueError):
        load_bundle(complete_path)


def test_stale_image(complete_path):
    manifest_edit(complete_path, lambda m: m["weather_images"][0].update(
        valid_start="2026-08-30T00:00:00+08:00", valid_end="2026-08-31T00:00:00+08:00"))
    with pytest.raises(ValueError, match="image outside"):
        load_bundle(complete_path)


def test_unknown_forecast_pollutants_need_reason(complete_path):
    p = complete_path.parent / "forecasts.json"
    data = json.loads(p.read_text())
    data[0]["primary_pollutants"] = []
    write_json(p, data)
    with pytest.raises(ValueError, match="partial forecast reason"):
        load_bundle(complete_path)
    data[0]["reason"] = "上游未提供首要污染物"
    write_json(p, data)
    assert any(r["section"] == "forecast/2026-09-01" for r in build_view(load_bundle(complete_path))["missing"])


def test_media_hash_order_and_no_hidden_text(tmp_path, complete):
    from airmore.regional.render import preflight, render_docx, verify_docx
    b = load_bundle(complete)
    view = build_view(b)
    target = tmp_path / "media.docx"
    media = render_docx(b, view, target, preflight())
    doc = Document(target)
    weather = doc.tables[-1]
    assert len(weather.rows) == 1 and len(weather.columns) == 2
    assert all(c.paragraphs[0].paragraph_format.keep_with_next for c in weather.rows[0].cells)
    notes = [p for c in weather.rows[0].cells for p in c.paragraphs if "有效时段：" in p.text]
    assert len(notes) == 2
    assert all(p.paragraph_format.keep_together for p in notes)
    from docx.oxml.ns import qn
    assert weather.rows[0]._tr.find(".//" + qn("w:cantSplit")) is not None
    assert doc.styles["Normal"].element.rPr.rFonts.get(qn("w:eastAsia")) == preflight()["docx_font_family"]
    assert all(not k.lower().endswith("theme") for name in ("Heading 1", "Heading 2")
               for k in doc.styles[name].element.rPr.rFonts.attrib)
    verify_docx(target, view, media)
    weather.rows[0].cells[0].add_paragraph("AQI 999")
    doc.save(target)
    with pytest.raises(ValueError, match="readback"):
        verify_docx(target, view, media)


@pytest.mark.parametrize("key,bad", [("sources", {}), ("stations", ["bad"]), ("files", []),
                                     ("region", "bad")])
def test_structural_types(bundle_path, key, bad):
    manifest_edit(bundle_path, lambda m: m.update({key: bad}))
    with pytest.raises(ValueError):
        load_bundle(bundle_path)


def test_concurrent_writer_lock(tmp_path, complete):
    from airmore.regional.pipeline import generate
    root = tmp_path / "out"
    root.mkdir()
    lock = root / ".示例区域_2026-09-01_第TEST001期.lock"
    lock.touch()
    with pytest.raises(ValueError, match="in progress"):
        generate(complete, root)
    assert lock.exists() and list(root.iterdir()) == [lock]


def test_advice_fact_binding(complete_path):
    p = complete_path.parent / "advice.json"
    data = json.loads(p.read_text())
    data[0].update(text="综合指数999.99，应立即管控。")
    write_json(p, data)
    with pytest.raises(ValueError, match="fact placeholders"):
        load_bundle(complete_path)
    data[0].update(text="参考{P01_month_composite_stat}，安排复核。", evidence_ids=["P01_month_composite_stat"])
    write_json(p, data)
    view = build_view(load_bundle(complete_path))
    assert any("综合指数 2.16" in x.get("text", "") for x in view["blocks"])
    data[0].update(entity_id="P03")
    write_json(p, data)
    with pytest.raises(ValueError, match="entity mismatch"):
        load_bundle(complete_path)


def test_complete_view_and_source_bound_charts(complete):
    b = load_bundle(complete)
    v = build_view(b)
    assert not v["missing"]
    for chart in v["charts"]:
        for panel in chart["panels"]:
            assert sum(panel["ratios"].values()) == pytest.approx(1)
            assert len(panel["refs"]) == 6
    table = next(x for x in v["blocks"] if x["type"] == "table")
    assert table["styles"][-1]["color"] == "FF0000"
    assert table["styles"][1]["color"] == "806000"



@pytest.mark.parametrize("count", [1, 4, 28, 50])
def test_real_docx_readback(tmp_path, count):
    from airmore.regional.render import preflight, render_docx, verify_docx
    b = load_bundle(make_example(tmp_path / "in", count=count, missing=True))
    v = build_view(b)
    target = tmp_path / "report.docx"
    media = render_docx(b, v, target, preflight())
    verify_docx(target, v, media)
    doc = Document(target)
    assert len(doc.tables[0].rows) == count + 1
    from docx.oxml.ns import qn
    assert doc.tables[0].rows[0]._tr.find(".//" + qn("w:tblHeader")) is not None
    doc.tables[0].rows[1].cells[-2].text = "999.99"
    doc.save(target)
    with pytest.raises(ValueError, match="readback"):
        verify_docx(target, v, media)


def test_atomic_export_and_nonoverwrite(tmp_path, complete):
    from airmore.regional.pipeline import generate
    result = generate(complete, tmp_path / "out")
    evidence = json.loads(Path(result["json"]).read_text())
    assert result["status"] == "complete"
    assert all(Path(result[k]).is_file() for k in ("docx", "pdf", "json"))
    for name, digest in evidence["artifact_sha256"].items():
        assert hashlib.sha256((Path(result["directory"]) / name).read_bytes()).hexdigest() == digest
    with pytest.raises(ValueError, match="already exists"):
        generate(complete, tmp_path / "out")


def test_failed_conversion_publishes_nothing(tmp_path, complete, monkeypatch):
    from airmore.regional import pipeline
    def fail(*args):
        raise RuntimeError("conversion failure")
    monkeypatch.setattr(pipeline, "convert_pdf", fail)
    with pytest.raises(RuntimeError, match="conversion failure"):
        pipeline.generate(complete, tmp_path / "out")
    assert not list((tmp_path / "out").iterdir())
