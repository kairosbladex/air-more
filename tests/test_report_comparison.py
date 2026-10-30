from copy import deepcopy
import io

from docx import Document
from PIL import Image
import pytest

from airmore.service.models import METRICS, UNITS
from airmore.service.report_comparison import build_comparison, change
from airmore.service.report_index import evaluate
from tests.test_hangar_trend_alignment import case, record, statistics
from tests.test_report_platform_index import survey


@pytest.mark.parametrize('previous,current,delta,percent,text', [
    (10,12.5,2.5,25,'上升25.00%'),(12.5,10,-2.5,-20,'下降20.00%'),
    (10,10,0,0,'持平'),(0,0,0,None,'持平'),(0,10,10,None,'比较期为0，未计环比'),
    (None,10,None,None,'比较期缺测'),(10,None,None,None,'本期缺测'),
    (None,None,None,None,'两期缺测'),(10000,10000.0001,.0001,.000001,'上升<0.01%'),
])
def test_independent_period_arithmetic(previous,current,delta,percent,text):
    result=change(previous,current)
    assert result['delta']==pytest.approx(delta) if delta is not None else result['delta'] is None
    assert result['percent']==pytest.approx(percent) if percent is not None else result['percent'] is None
    assert result['change_text']==text


def comparison_stats():
    req,data=case('WEEK',current=[record('2026-09-14')],raw=False)
    stats=statistics(req,data);h=stats['hangars']['h'];h.pop('indicator_validity',None)
    h['previous']['survey']=survey(pm10=20,pm2_5=10)
    h['current']['survey']=survey(pm10=10,pm2_5=12.5)
    return stats


def test_units_missing_zero_and_dynamic_indicator_catalog():
    stats=comparison_stats();before=deepcopy(stats)
    result=build_comparison(stats,evaluate(stats))
    assert [r['change_text'] for r in result['rows'][1:]]==['下降50.00%','上升25.00%']
    assert 'PM10' in result['leads']['particles'] and '下降50.00%' in result['leads']['particles']
    assert stats==before
    h=stats['hangars']['h'];h['previous']['survey']['pm10']['unit']='mg/m³'
    assert build_comparison(stats,evaluate(stats))['rows'][1]['change_text']=='单位不一致'
    h['indicator_validity']={m:{'valid':True,'unit':UNITS[m]} for m in METRICS}
    h['indicator_validity']['pm10']['valid']=False
    row=build_comparison(stats,evaluate(stats))['rows'][1]
    assert row['previous_text']==row['current_text']=='/' and row['change_text']=='单位无效'
    stats['indicator_catalog']['table_metrics']=['so2','pm2_5']
    assert [r['metric'] for r in build_comparison(stats,evaluate(stats))['rows']]==['index','pm2_5']


def test_panels_follow_current_rank_and_titles_share_picture_cells(tmp_path,monkeypatch):
    from airmore.service import approved_report as report
    from airmore.service.approved_advice import write_report
    from airmore.service.config import Settings
    stats=comparison_stats();proto=stats['hangars'].pop('h')
    # Insertion order differs from period rank; the last-day rank also differs.
    for hid,name,current,last in [('a','甲机库',70,1),('b','乙机库',None,60),('c','丙机库',60,70)]:
        h=deepcopy(proto);h['name']=name
        h['current']['survey']=survey(pm2_5=current) if current is not None else {m:{'value':None,'unit':UNITS[m]} for m in METRICS}
        h['days'][stats['request']['end_date']]['survey']=survey(pm2_5=last)
        stats['hangars'][hid]=h
    index=evaluate(stats);before=deepcopy(stats);captured=[]
    picture=io.BytesIO();Image.new('RGB',(20,20),'white').save(picture,format='PNG')
    monkeypatch.setattr(report,'pie',lambda *args:picture.getvalue())
    monkeypatch.setattr(report,'forecast_card',lambda *args:picture.getvalue())
    def trend(h,*args):captured.append(h['name']);return picture.getvalue()
    monkeypatch.setattr(report,'measured_trend',trend)
    result=report.render_document(stats,write_report(None,stats),index,Settings(),tmp_path,{'font_family':'DejaVu Sans'})
    doc=Document(tmp_path/'report.docx');period_names=[r.cells[0].text for r in doc.tables[2].rows[1:]]
    assert period_names==['丙机库','甲机库','乙机库']
    assert [r.cells[0].text for r in doc.tables[1].rows[1:]]!=period_names
    assert captured==period_names==[r.cells[0].text for r in doc.tables[4].rows[1:]]
    pies=doc.tables[3]
    assert all(t.alignment==1 for t in doc.tables[1:])
    assert [pies.rows[i].cells[1].text for i in (0,2,4)]==period_names
    for row in pies.rows[1::2]:
        pair=row.cells[0].tables[0]
        assert pair.alignment==1
        assert [c.paragraphs[0].text for c in pair.rows[0].cells]==['比较期占比','本期占比']
        for cell in pair.rows[0].cells:
            assert len(cell._tc.xpath('.//w:drawing'))==1
            assert all(p.alignment==1 for p in cell.paragraphs)
    assert doc.tables[-1].cell(0,4).text=='变化'
    assert '甲机库的PM2.5为70.00' in '\n'.join(p.text for p in doc.paragraphs)
    assert '上升600.00%' in '\n'.join(p.text for p in doc.paragraphs)
    assert result['hangar_display_order']==['c','a','b'] and stats==before


def test_action_facts_preserve_direction_while_numbers_stay_program_bound():
    import re
    from airmore.service.approved_advice import facts_for, write_report
    stats=comparison_stats();stats['hangars']['h']['name']='testjk122'
    stats['report_index']=evaluate(stats)
    facts=facts_for(stats)
    assert '可吸入颗粒物' in facts['particles'] and '下降' in facts['particles']
    assert not re.search(r'[0-9%]', facts['particles']+facts['comparison'])
    audit=write_report(None,stats)['h'].metadata['report_advice']
    row=next(r for r in audit['period_comparison']['rows'] if r['metric']=='pm10')
    assert row['hangar_name']=='testjk122'
    assert (row['previous'],row['current'],row['percent'])==(20,10,-50)
    assert 'testjk122' in audit['period_comparison']['leads']['particles']
    assert '下降50.00%' in audit['period_comparison']['leads']['particles']
